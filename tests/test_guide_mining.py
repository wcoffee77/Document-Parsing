"""guide_mining.probe / stats — 합성 docx로 '실효 서식'이 제대로 풀리는지 확인한다.

정답은 우리가 docx를 만들 때 넣은 값이다(스타일 상속·직접 서식·번호 정의 각각).
"""

from __future__ import annotations

import json

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt
from typer.testing import CliRunner

from doc2report.cli import app
from doc2report.guide_mining import collect_docx, probe_all, write_outputs
from doc2report.guide_mining.probe import _typed_marker, probe_docx
from doc2report.guide_mining.stats import _ending_class, summarize


def _make(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(15)
    sec.left_margin, sec.right_margin = Mm(25), Mm(20)
    # 스타일 상속: Normal → 본문 크기 13pt, 줄간격 1.5
    normal = doc.styles["Normal"]
    normal.font.size = Pt(13)
    normal.element.get_or_add_rPr().append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="바탕체"/>'))
    normal.paragraph_format.line_spacing = 1.5
    doc.add_paragraph("□ 추진 배경을 정리함")                       # 글자 말머리 + 스타일 상속
    doc.add_paragraph("")                                              # 빈 문단
    p = doc.add_paragraph("- 세부 항목")                              # 직접 서식이 스타일을 덮음
    p.paragraph_format.left_indent = Mm(8)
    p.paragraph_format.space_after = Pt(6)
    p.runs[0].font.size = Pt(11)
    p.runs[0].bold = True
    doc.add_paragraph("​")                                        # 보이지 않는 글자만 든 줄
    doc.add_paragraph("-5%p 개선됨").alignment = WD_ALIGN_PARAGRAPH.RIGHT  # 말머리 아님
    doc.add_paragraph("2026. 9. 30. 기준 매출 1,200억원 (▲5%)")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "구분"
    table.cell(0, 1).text = "내용"
    table.cell(1, 0).text = "현황"
    table.cell(1, 1).text = "양호함"
    trpr = table.rows[0]._tr.get_or_add_trPr()
    trpr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    table.cell(0, 0)._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="F2F2F2"/>'))
    # 머리말·꼬리말(쪽번호 필드)
    sec.header.paragraphs[0].text = "대외비"
    footer = sec.footer.paragraphs[0]
    footer._p.append(parse_xml(
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'))
    footer._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:instrText> PAGE </w:instrText></w:r>'))
    footer._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'))
    # 텍스트 상자(결재란): 본문 문단으로 새면 안 되고 따로 세어야 한다
    box = parse_xml(
        f'<w:p {nsdecls("w")}><w:r><w:pict><w:txbxContent>'
        '<w:p><w:r><w:t>결재</w:t></w:r></w:p></w:txbxContent></w:pict></w:r></w:p>')
    doc.element.body.insert(0, box)
    doc.save(path)


@pytest.fixture()
def sample(tmp_path):
    path = tmp_path / "정식보고서.docx"
    _make(path)
    return path


def test_page_and_style_inheritance(sample):
    probe = probe_docx(sample)
    page = probe.sections[0]
    assert (page.top_mm, page.bottom_mm, page.left_mm, page.right_mm) == (20.0, 15.0, 25.0, 20.0)
    first = next(p for p in probe.paragraphs if p.text.startswith("□"))
    assert first.fmt.size_pt == 13 and first.fmt.east_asia == "바탕체"   # 스타일에서 상속
    assert first.line_rule == "auto" and first.line_value == pytest.approx(1.5)
    direct = next(p for p in probe.paragraphs if p.text.startswith("- 세부"))
    assert direct.fmt.size_pt == 11 and direct.fmt.bold is True          # 직접 서식이 이김
    assert direct.left_mm == pytest.approx(8, abs=0.1)
    assert direct.space_after_pt == 6
    assert direct.fmt.east_asia == "바탕체"                                # 안 덮은 값은 계속 상속


def test_blank_lines_are_kept_not_dropped(sample):
    body = [p for p in probe_docx(sample).paragraphs if p.where == "body"]
    blanks = [p for p in body if p.blank]
    assert len(blanks) == 2                       # 진짜 빈 줄 + U+200B 줄 (docx_reader는 둘 다 버린다)
    assert any(p.invisible_codes == ["U+200B"] for p in blanks)


def test_markers(sample):
    body = {p.text: p for p in probe_docx(sample).paragraphs if p.where == "body"}
    assert body["□ 추진 배경을 정리함"].marker == "□"
    assert body["□ 추진 배경을 정리함"].marker_kind == "typed"
    assert body["□ 추진 배경을 정리함"].marker_sep == "space"
    assert body["- 세부 항목"].marker == "-"
    assert body["-5%p 개선됨"].marker_kind is None       # 공백 없이 바로 숫자 = 값 표기


@pytest.mark.parametrize("text,expected", [
    ("ㆍ입사예정", ("ㆍ", "none")),        # 기호는 붙여 써도 말머리
    ("1.\t현황", ("1.", "tab")),
    ("1.5배 증가", None),
    ("ㅇㅇ팀 소속", None),
    ("○○팀 소속", None),                  # 같은 기호 반복 = 자리표시자
    ("(1) 세부", ("(1)", "space")),
    ("가. 세부", ("가.", "space")),
    ("  ", None),
])
def test_typed_marker(text, expected):
    assert _typed_marker(text) == expected


def test_table_header_footer_textbox(sample):
    probe = probe_docx(sample)
    table = probe.tables[0]
    assert (table.rows, table.cols, table.header_rows, table.header_fill) == (2, 2, 1, "F2F2F2")
    assert any(p.where == "table" and p.text == "양호함" for p in probe.paragraphs)
    header = next(h for h in probe.headers_footers if h.kind == "header")
    footer = next(h for h in probe.headers_footers if h.kind == "footer")
    assert header.text == "대외비" and not header.has_page_field
    assert footer.has_page_field
    assert probe.textbox_count == 1
    assert [p.text for p in probe.paragraphs if p.where == "textbox"] == ["결재"]
    assert all("결재" not in p.text for p in probe.paragraphs if p.where == "body")


def test_summary_and_privacy(sample, tmp_path):
    probes, skipped = probe_all(collect_docx([sample.parent]))
    assert not skipped and len(probes) == 1
    summary = summarize(probes)
    assert summary["page"]["layout"][0]["value"].startswith("A4 portrait 위20.0 아래15.0")
    assert summary["blank_lines"]["invisible_chars"] == {"U+200B": 1}
    assert summary["notation"]["date_forms"]["YYYY.M.D."] == 1
    assert summary["notation"]["number_forms"]["억원"] == 1
    assert summary["notation"]["number_forms"]["▲/△ 증감"] == 1
    assert summary["headers_footers"]["footer"]["page_field_share"] == 1.0
    out = write_outputs(probes, skipped, tmp_path / "out")
    dumped = (out / "probe" / "doc001.json").read_text(encoding="utf-8")
    assert "추진 배경" not in dumped and "정식보고서" not in dumped      # 기본은 원문·파일 이름을 뺀다
    assert "추진 배경" not in (out / "probe_summary.json").read_text(encoding="utf-8")
    assert "# 정식보고서 말뭉치" in (out / "probe_summary.md").read_text(encoding="utf-8-sig")
    kept = write_outputs(probes, skipped, tmp_path / "out2", keep_text=True, keep_names=True)
    assert "추진 배경" in (kept / "probe" / "doc001.json").read_text(encoding="utf-8")


@pytest.mark.parametrize("sentence,expected", [
    ("성과를 개선하였습니다.", "~습니다 (합쇼체)"),
    ("성과 개선함", "~음/함/임 (개조식)"),
    ("성과 개선 완료", "명사 종결"),
    ("성과가 개선되었다.", "~다 (서술)"),
    ("목표 달성률 95%", "수치·기호 종결"),
])
def test_ending_class(sentence, expected):
    assert _ending_class(sentence) == expected


def test_marker_layers_have_endings_and_are_ordered_by_indent(sample):
    rows = summarize([probe_docx(sample)])["markers"]["by_marker"]
    lefts = [r["left_mm_median"] for r in rows]
    assert lefts == sorted(lefts)                      # 얕은 층부터
    box = next(r for r in rows if r["marker"] == "글자 □")
    assert box["ending_class"][0]["value"] == "~음/함/임 (개조식)"   # "정리함"
    assert box["period_ended_share"] == 0.0


def test_drm_file_is_skipped_with_reason(tmp_path):
    fake = tmp_path / "drm.docx"
    fake.write_bytes(b"\x00DRM-ENCRYPTED")
    probes, skipped = probe_all([fake])
    assert not probes and "DRM" in skipped[0]


def test_cli_probe(sample, tmp_path):
    result = CliRunner().invoke(app, ["probe", str(sample.parent), "-o", str(tmp_path / "res")])
    assert result.exit_code == 0, result.output
    assert "1건 분석" in result.output
    assert json.loads((tmp_path / "res" / "probe_summary.json").read_text(encoding="utf-8"))["docs"]["count"] == 1


_WPS_NS = 'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
_NOTE_RUN = (
    '<w:p><w:pPr><w:jc w:val="left"/></w:pPr><w:r><w:rPr><w:rFonts w:eastAsia="바탕체"/>'
    '<w:color w:val="0000ff"/><w:sz w:val="20"/></w:rPr>'
    '<w:t>※ 최근 평가 상위, 영어회화 2급 이상</w:t></w:r></w:p>')


def _note_doc(path, *, vml: bool):
    """본문 문장 아래·오른쪽 여백에 붙은 주석 상자(바탕체 10pt, 파란 글씨)."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    p = doc.add_paragraph("□ 핵심인력 선정")
    if vml:
        shape = (
            f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"><w:pict>'
            '<v:shape style="position:absolute;margin-left:10pt;margin-top:8pt;width:85pt;height:30pt;'
            'mso-position-horizontal-relative:margin;mso-position-vertical-relative:text" '
            'stroked="f" filled="f"><v:textbox><w:txbxContent>' + _NOTE_RUN +
            '</w:txbxContent></v:textbox></v:shape></w:pict></w:r>')
    else:
        shape = (
            f'<w:r {nsdecls("w", "wp", "a")} {_WPS_NS}><w:drawing><wp:anchor behindDoc="0">'
            '<wp:positionH relativeFrom="page"><wp:posOffset>7020000</wp:posOffset></wp:positionH>'
            '<wp:positionV relativeFrom="paragraph"><wp:posOffset>216000</wp:posOffset></wp:positionV>'
            '<wp:extent cx="540000" cy="360000"/><wp:wrapSquare wrapText="bothSides"/>'
            '<a:graphic><a:graphicData><wps:wsp><wps:spPr><a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr>'
            '<wps:txbx><w:txbxContent>' + _NOTE_RUN +
            '</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic>'
            '</wp:anchor></w:drawing></w:r>')
    p._p.append(parse_xml(shape))
    doc.save(path)


def test_annotation_box_drawingml(tmp_path):
    path = tmp_path / "note.docx"
    _note_doc(path, vml=False)
    probe = probe_docx(path)
    box = probe.textboxes[0]
    assert box.kind == "drawingml" and box.floating
    assert box.placement == "오른쪽 여백 / 문단 기준 아래"
    assert (box.width_mm, box.height_mm, box.wrap) == (15.0, 10.0, "wrapSquare")
    assert (box.border, box.fill) == ("none", "none")
    assert (box.anchor_where, box.anchor_marker) == ("body", "□")
    note = next(p for p in probe.paragraphs if p.where == "textbox")
    assert (note.fmt.east_asia, note.fmt.size_pt, note.fmt.color) == ("바탕체", 10.0, "0000FF")
    assert note.box == 0
    ann = summarize([probe])["annotations"]
    assert ann["color"][0]["value"] == "0000FF" and ann["size_pt"][0]["value"] == 10.0
    assert ann["anchor_marker"][0]["value"] == "□"
    assert json.dumps(probe.to_dict())  # 글 없이도 직렬화된다
    assert "영어회화" not in json.dumps(probe.to_dict(), ensure_ascii=False)


def test_annotation_box_vml(tmp_path):
    path = tmp_path / "note_vml.docx"
    _note_doc(path, vml=True)
    box = probe_docx(path).textboxes[0]
    assert box.kind == "vml" and box.floating
    assert box.border == "none" and box.fill == "none"
    assert box.v_rel == "text" and box.placement.endswith("문단 기준 아래")
    assert box.width_mm == pytest.approx(30.0, abs=0.1)   # 85pt


def test_layers_with_equal_indent_are_ordered_by_leading_spaces(tmp_path):
    path = tmp_path / "spaces.docx"
    doc = Document()
    for text in ("   - 세부", " □ 큰항목", "□ 다른항목", "   - 세부2"):
        doc.add_paragraph(text)
    doc.save(path)
    rows = summarize([probe_docx(path)])["markers"]["by_marker"]
    assert [r["marker"] for r in rows] == ["글자 □", "글자 -"]


def _rich_doc(path, heading="1. 추진 배경"):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    doc.add_paragraph("보고서 제목").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("2026. 9. 30.").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph(heading)
    p = doc.add_paragraph()
    p.add_run("□ ")
    p.add_run("핵심인력: ").bold = True          # 말머리 뒤 앞부분만 굵게
    p.add_run("조건을 충족한 인원")
    doc.add_paragraph("【현황】")
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate(("구분", "내용"), start=0):
        table.cell(0, r).text = row
    table.cell(1, 0).text = "현황"
    table.cell(1, 1).text = "양호함"
    for r in (0, 1):                                 # 첫 열만 음영
        table.cell(r, 0)._tc.get_or_add_tcPr().append(
            parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="D9D9D9"/>'))
    doc.add_paragraph("※ 표 주석")
    doc.save(path)


def test_bold_pattern_table_context_and_shape(tmp_path):
    path = tmp_path / "rich.docx"
    _rich_doc(path)
    probe = probe_docx(path)
    item = next(x for x in probe.paragraphs if x.marker == "□")
    assert item.bold_pattern == "prefix"
    table = probe.tables[0]
    assert (table.before_kind, table.after_kind) == ("꺾쇠 제목", "※ 주석")
    assert table.fill_first_col_share == 1.0 and table.fill_first_row_share == 0.5   # 첫 열은 1행 이후만 센다
    assert table.fill_other_share == 0.0 and table.fill_colors == ["D9D9D9"]
    cells = {x.text: (x.cell_row, x.cell_col) for x in probe.paragraphs if x.where == "table"}
    assert cells["구분"] == (0, 0) and cells["양호함"] == (1, 1)
    s = summarize([probe])
    assert s["dates"]["shapes"][0]["value"] == "9999. 9. 99."
    assert s["skeleton"]["first"][0]["shapes"][0]["value"].startswith("center·")
    assert any("날짜형" in r["value"] for r in s["skeleton"]["first"][1]["shapes"])
    assert s["table_cells"]["첫 열(1행 이후)"]["count"] == 1


def test_phrases_only_repeated_across_docs_and_optional(tmp_path):
    a, b = tmp_path / "a.docx", tmp_path / "b.docx"
    _rich_doc(a)
    _rich_doc(b, heading="2. 비밀 사업명")
    s = summarize([probe_docx(a), probe_docx(b)])
    titles = {r["value"]: r["docs"] for r in s["phrases"]["section_titles"]}
    assert "비밀 사업명" not in titles                       # 한 문서에만 있는 말은 싣지 않는다
    assert {r["value"] for r in s["phrases"]["table_terms"]} >= {"구분", "현황"}
    assert summarize([probe_docx(a)], phrases=False)["phrases"] is None


def test_full_markdown_renders_all_new_sections(tmp_path):
    a = tmp_path / "docs" / "a.docx"
    a.parent.mkdir()
    _rich_doc(a)
    _note_doc(tmp_path / "docs" / "n.docx", vml=False)
    result = CliRunner().invoke(app, ["probe", str(a.parent), "-o", str(tmp_path / "res")])
    assert result.exit_code == 0, result.output
    md = (tmp_path / "res" / "probe_summary.md").read_text(encoding="utf-8-sig")
    for heading in ("말머리 없는 줄", "문서 첫머리·말미 구성", "주석 상자", "표 칸 서식", "여러 문서에 반복되는 말"):
        assert heading in md
    assert "핵심인력" not in md and "영어회화" not in md      # 원문 글자는 새지 않는다
    off = CliRunner().invoke(app, ["probe", str(a.parent), "-o", str(tmp_path / "res2"), "--no-phrases"])
    assert "여러 문서에 반복되는 말" not in (tmp_path / "res2" / "probe_summary.md").read_text(encoding="utf-8-sig")


def test_theme_shading_children_bold_and_charset(tmp_path):
    path = tmp_path / "theme.docx"
    doc = Document()
    parent = doc.add_paragraph()
    parent.add_run("□ ").bold = False
    parent.add_run("하위 있는 항목").bold = True
    doc.add_paragraph("   - 세부 내용 無 ▲5 ↑")
    leaf = doc.add_paragraph()
    leaf.add_run("□ ")
    leaf.add_run("하위 없는 항목")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "구분 BT"
    table.cell(1, 0).text = "값"
    table.cell(0, 0)._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="auto" '
        'w:themeFill="background1" w:themeFillShade="F2"/>'))
    table.cell(0, 0)._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="8"/><w:top w:val="nil"/></w:tcBorders>'))
    doc.save(path)
    probe = probe_docx(path)
    t = probe.tables[0]
    assert t.header_fill == "theme:background1+shadeF2"      # fill=auto여도 테마 음영을 놓치지 않는다
    assert t.fill_first_row_share == 1.0
    assert t.cell_border_kinds == ["single/1pt"] and t.cell_border_share == 0.5   # nil은 테두리 아님
    s = summarize([probe_docx(path, doc_id="a"), probe_docx(path, doc_id="b")])
    by = next(r for r in s["markers"]["by_marker"] if r["marker"] == "글자 □")["bold_by_children"]
    assert by["하위 항목 있음"][0]["value"] == "all" and by["하위 항목 없음"][0]["value"] == "none"
    assert {e["value"] for e in s["charset"]["symbols"]} >= {"▲", "↑"}
    assert s["charset"]["hanja"][0]["value"] == "無"
    assert "BT" in {r["value"] for r in s["phrases"]["latin_terms"]}
    assert s["docs"]["per_doc"][0]["margins_mm"] is not None


def test_leading_space_distribution_and_fullwidth(tmp_path):
    path = tmp_path / "lead.docx"
    doc = Document()
    for text in ("  □ 하나", "  □ 둘", "\u3000□ 셋", "    - 넷", "    - 다섯", "    - 여섯"):
        doc.add_paragraph(text)
    doc.save(path)
    rows = {r["marker"]: r for r in summarize([probe_docx(path)])["markers"]["by_marker"]}
    assert rows["글자 □"]["leading_spaces"][0] == {"value": 2, "count": 2, "share": 0.667}
    assert rows["글자 -"]["leading_spaces"][0]["value"] == 4
    assert rows["글자 □"]["leading_wide_share"] == 0.333 and rows["글자 -"]["leading_wide_share"] == 0.0
