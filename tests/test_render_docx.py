"""렌더링 결과가 실제 docx XML에 반영되는지 — 서식이 바뀔 때 회귀를 잡는 그물."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.oxml.ns import qn

from doc2report.pipeline import convert
from doc2report.profile import load_profile
from doc2report.render.base_template import open_base_template
from doc2report.render.docx_writer import _base_template
from doc2report.units import emu_to_dxa, emu_to_pt

FIXTURE = Path(__file__).parent / "fixtures" / "sample_report.md"


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("render") / "report.docx"
    result = convert(str(FIXTURE), out, "default")
    return result, DocxDocument(str(out))


def test_document_opens_and_has_content(built):
    _, docx = built
    assert len(docx.paragraphs) > 5
    assert len(docx.tables) == 1


def test_header_row_is_shaded_with_profile_color(built):
    """표 머리행 음영은 profiles의 tables.header_shading을 그대로 써야 한다
    (2026-09-29 사용자 요청: R242,G242,B242 옅은 회색 = F2F2F2)."""
    _, docx = built
    profile = load_profile("default")
    header_cell = docx.tables[0].rows[0].cells[0]
    shd = header_cell._tc.tcPr.find(qn("w:shd"))
    assert shd is not None
    assert shd.get(qn("w:fill")).upper() == profile.tables.header_shading


def test_base_template_is_embedded_not_a_docx_file():
    """기본 템플릿을 .docx 파일로 두면 사내 PC 문서보안/백신이 손상시켜 "Package not
    found"로 죽었다(2026-09-29, 두 번). 파이썬 소스에 든 사본을 메모리에서 연다."""
    import zipfile

    assert not (Path(__file__).parents[1] / "assets" / "default_template.docx").exists()
    profile = load_profile("default")
    assert profile.template is None
    source = _base_template(profile)
    assert zipfile.is_zipfile(source)
    source.seek(0)
    assert len(DocxDocument(source).styles) > 10
    assert open_base_template().getvalue()[:2] == b"PK"


def test_page_setup_matches_profile(built):
    _, docx = built
    profile = load_profile("default")
    section = docx.sections[0]
    # docx는 길이를 twips로 저장하므로 1 twip(635 EMU) 범위의 반올림 오차를 허용한다.
    assert abs(int(section.page_width) - profile.page.width) <= 635
    assert abs(int(section.top_margin) - profile.page.margin.top) <= 635
    assert abs(int(section.left_margin) - profile.page.margin.left) <= 635


def test_table_never_exceeds_usable_width(built):
    result, docx = built
    profile = load_profile("default")
    limit = emu_to_dxa(profile.page.usable_width)
    for table in docx.tables:
        widths = [int(col.get(qn("w:w")))
                  for col in table._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))]
        assert sum(widths) <= limit, f"표 폭 {sum(widths)} > 사용가능폭 {limit}"


def test_table_layout_is_fixed_and_header_repeats(built):
    _, docx = built
    tbl_pr = docx.tables[0]._tbl.tblPr
    assert tbl_pr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
    first_row = docx.tables[0].rows[0]._tr
    assert first_row.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None


def test_korean_font_is_written_to_east_asia_attribute(built):
    _, docx = built
    profile = load_profile("default")
    body = profile.font("body")
    found = set()
    for paragraph in docx.paragraphs:
        for run in paragraph.runs:
            rpr = run._element.find(qn("w:rPr"))
            if rpr is None:
                continue
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None and fonts.get(qn("w:eastAsia")):
                found.add(fonts.get(qn("w:eastAsia")))
    assert body.east_asia in found


def test_body_font_size_matches_profile(built):
    _, docx = built
    profile = load_profile("default")
    expected = emu_to_pt(profile.font("body").size)
    sizes = [emu_to_pt(int(run.font.size))
             for paragraph in docx.paragraphs for run in paragraph.runs
             if run.font.size and paragraph.text.strip()]
    assert expected in sizes


def test_list_markers_come_from_profile(built):
    _, docx = built
    profile = load_profile("default")
    markers = {level.marker for level in profile.numbering if level.marker}
    texts = [p.text for p in docx.paragraphs]
    assert any(any(text.startswith(m) for m in markers) for text in texts)


def test_changes_are_reported(built):
    result, _ = built
    assert result.changes
    assert "개조식" in result.report()


def _render_table(tmp_path, cell_blocks):
    from doc2report.ir import Cell, Document, Row, Table
    from doc2report.render.docx_writer import DocxRenderer

    profile = load_profile("confluence")
    table = Table(rows=[Row(cells=[Cell(blocks=cell_blocks)])], header_rows=0)
    out = tmp_path / "t.docx"
    DocxRenderer(profile).save(Document(blocks=[table]), out)
    return DocxDocument(str(out)).tables[0].rows[0].cells[0], profile


def test_heading_inside_cell_is_bold_table_size_without_blank_line_above(tmp_path):
    """표 안 굵은 제목이 14pt로 커지고 위에 빈 줄이 생기던 문제(2026-09-29 사용자)."""
    from doc2report.ir import Heading, Run

    cell, profile = _render_table(tmp_path, [Heading(level=3, runs=[Run("굵은 소제목")])])
    assert [p.text for p in cell.paragraphs] == ["굵은 소제목"]  # 앞에 빈 문단 없음
    run = cell.paragraphs[0].runs[0]
    assert run.bold
    assert emu_to_pt(run.font.size) == emu_to_pt(profile.font("table").size)


def test_list_item_inside_cell_uses_table_font_and_cell_marker(tmp_path):
    from doc2report.ir import ListItem, Run

    cell, profile = _render_table(tmp_path, [ListItem(depth=0, runs=[Run("항목")])])
    assert [p.text for p in cell.paragraphs] == ["- 항목"]
    assert emu_to_pt(cell.paragraphs[0].runs[0].font.size) == emu_to_pt(profile.font("table").size)


def test_reference_mark_is_two_points_smaller_than_body(tmp_path):
    """※ 참고사항은 항상 본문보다 2pt 작게(2026-09-29 사용자)."""
    from doc2report.ir import Document, ListItem, Paragraph, Run
    from doc2report.render.docx_writer import DocxRenderer

    profile = load_profile("confluence")  # 본문 12pt
    doc = Document(blocks=[Paragraph(runs=[Run("일반 문단")]),
                           Paragraph(runs=[Run("※ 참고 문단")]),
                           ListItem(depth=1, runs=[Run("참고 항목")], marker="※")])
    out = tmp_path / "n.docx"
    DocxRenderer(profile).save(doc, out)
    sizes = {p.text: emu_to_pt(p.runs[-1].font.size) for p in DocxDocument(str(out)).paragraphs if p.text}
    assert sizes["일반 문단"] == 12
    assert sizes["※ 참고 문단"] == 10
    assert sizes["※\t참고 항목"] == 10


def test_note_size_delta_survives_dump_and_reload(tmp_path):
    from doc2report.profile import dump_profile

    profile = load_profile("default")
    path = tmp_path / "p.yaml"
    path.write_text(dump_profile(profile), encoding="utf-8")
    assert load_profile(path).text.note_size_delta == profile.text.note_size_delta


def test_table_itself_is_right_aligned(built):
    """표 전체는 오른쪽 정렬(2026-09-29 사용자). 셀 안 글자 정렬과는 별개다."""
    from docx.enum.table import WD_TABLE_ALIGNMENT

    _, docx = built
    assert docx.tables[0].alignment == WD_TABLE_ALIGNMENT.RIGHT
