"""표는 한 쪽 안에서 보인다 — 쪽 사이에서 잘리지 않게(2026-10-08 사용자, 정식보고서 항상 적용)."""

import pytest
from docx import Document as OpenDocx

from doc2report.ir import Cell, Document, Heading, ListItem, Paragraph, Row, Run, Table
from doc2report.layout.flow import plan_flow
from doc2report.layout.table_fit import plan_tables
from doc2report.profile import load_profile
from doc2report.render import docx_writer


def _table() -> Table:
    def cell(text, header=False):
        return Cell(blocks=[Paragraph(runs=[Run(text)])], is_header=header)
    rows = [Row(cells=[cell("구분", True), cell("(1안) 외부", True), cell("(2안) 사내", True)])]
    for name in ("보안", "비용", "구축 기간", "기술 수준", "검토 의견"):
        rows.append(Row(cells=[cell(name), cell("- 즉시 사용 가능"), cell("- 약 6개월 소요")]))
    return Table(rows=rows)


def _doc(n: int) -> Document:
    blocks = [Heading(level=2, runs=[Run("배경")])]
    blocks += [ListItem(depth=1, runs=[Run(f"항목{k} : 재고는 180개이고 안전재고 대비 2.1주분이며 월평균 소요량은 360개")], marker="□")
               for k in range(n)]
    blocks += [Paragraph(runs=[Run("표 앞 도입 문장")]), _table(), Paragraph(runs=[Run("마무리 문장")])]
    return Document(blocks=blocks)


def _plan(n):
    prof = load_profile("formal")
    doc = _doc(n)
    layouts = plan_tables(doc, prof)
    return prof, doc, layouts, plan_flow(doc, prof, layouts)


def test_overflowing_table_pulls_line_spacing_in_slightly():
    found = None
    for n in range(5, 80):
        prof, doc, layouts, flow = _plan(n)
        if flow.line_scale:
            found = (n, prof, doc, layouts, flow)
            break
    assert found, "줄간격을 줄여 표를 한 쪽에 넣을 수 있는 길이가 없음"
    n, prof, doc, layouts, flow = found
    assert all(prof.tables.page_fit_min_scale <= s < 1.0 for s in flow.line_scale.values())
    assert any("한 쪽에 넣으려고" in note for note in flow.notes)
    assert not flow.relaxed                                  # 넉넉한 간격이 켜져 있으면 표를 밀어내므로 끈다


def test_scaled_blocks_get_smaller_line_spacing_in_the_docx(tmp_path):
    for n in range(5, 80):
        prof, doc, layouts, flow = _plan(n)
        if flow.line_scale:
            break
    else:
        pytest.fail("시험 문서를 못 만듦")
    out = tmp_path / "o.docx"
    docx_writer.DocxRenderer(prof, layouts, flow=flow).save(doc, out)
    spacings = [p.paragraph_format.line_spacing for p in OpenDocx(str(out)).paragraphs if "항목" in p.text]
    base = prof.font("body").line_spacing
    assert spacings and all(isinstance(s, float) and s < base for s in spacings), (spacings, base)


def test_table_rows_keep_with_next_and_lead_in_stays_with_table(tmp_path):
    prof, doc, layouts, flow = _plan(3)
    out = tmp_path / "o.docx"
    docx_writer.DocxRenderer(prof, layouts, flow=flow).save(doc, out)
    docx = OpenDocx(str(out))
    table = docx.tables[0]
    for row in table.rows[:-1]:
        assert all(p.paragraph_format.keep_with_next for c in row.cells for p in c.paragraphs)
    assert not any(p.paragraph_format.keep_with_next for c in table.rows[-1].cells for p in c.paragraphs)
    lead = next(p for p in docx.paragraphs if p.text.strip() == "표 앞 도입 문장")
    assert lead.paragraph_format.keep_with_next                    # 도입 문장은 표와 함께


def test_table_that_cannot_fit_after_squeezing_is_pushed_and_reported():
    notes = []
    for n in range(5, 120):
        _, _, _, flow = _plan(n)
        notes += [x for x in flow.notes if "다음 쪽으로 넘김" in x]
    assert notes                                                # 줄여도 안 되는 길이에서는 표를 통째로 넘기고 알린다


def test_other_profiles_are_unchanged(tmp_path):
    prof = load_profile("default")
    assert not prof.tables.keep_on_page
    doc = _doc(10)
    flow = plan_flow(doc, prof, plan_tables(doc, prof))
    assert flow.line_scale == {}
