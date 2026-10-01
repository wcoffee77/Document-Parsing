"""웹 화면용 입력 확장 — Word(.docx) 읽기, 붙여넣은 글, 여러 입력 합치기, Markdown/PDF 출력."""

from __future__ import annotations

import subprocess

import pytest
from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from doc2report.ir import Heading, ListItem, Paragraph, Table, plain
from doc2report.parsers.docx_reader import parse_docx
from doc2report.parsers.markdown import parse_markdown
from doc2report.parsers.plaintext import looks_like_markdown, text_to_markdown
from doc2report.pipeline import convert_many
from doc2report.profile import load_profile
from doc2report.render.markdown_writer import render_markdown, write_markdown
from doc2report.sources import load_source, load_text


# ── Word 읽기 ───────────────────────────────────────────────────────────


@pytest.fixture
def word_file(tmp_path):
    d = DocxDocument()
    d.add_paragraph("분기 보고", style="Title")
    d.add_heading("추진 배경", level=1)
    p = d.add_paragraph("일반 ")
    p.add_run("굵게").bold = True
    p.add_run(" 끝")
    d.add_paragraph("첫째", style="List Number")
    d.add_paragraph("둘째", style="List Number")
    d.add_paragraph("□ 원문에 친 말머리")
    t = d.add_table(rows=3, cols=3)
    for i, text in enumerate(["구분", "값", "비고"]):
        cell = t.cell(0, i)
        cell.text = text
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9D9D9")
        cell._tc.get_or_add_tcPr().append(shd)
    t.cell(1, 0).merge(t.cell(2, 0)).text = "세로병합"
    t.cell(1, 1).merge(t.cell(1, 2)).text = "가로병합"
    t.cell(2, 1).text = "x"
    t.cell(2, 2).text = "y"
    path = tmp_path / "입력.docx"
    d.save(str(path))
    return path


def test_docx_structure_is_read_in_order(word_file, tmp_path):
    doc = parse_docx(word_file, image_dir=tmp_path).document
    assert doc.title == "분기 보고"
    heading, para, first, second, marked, table = doc.blocks
    assert isinstance(heading, Heading) and heading.level == 2  # Word 제목1 = 첫 절 = Markdown ##
    assert [r.bold for r in para.runs] == [False, True, False]
    assert (first.marker, second.marker, first.ordered) == ("1.", "2.", True)  # 화면에 보이던 번호 그대로
    assert isinstance(marked, Paragraph) and plain(marked.runs).startswith("□")
    assert isinstance(table, Table) and table.header_rows == 1  # 음영 있는 첫 행 = 머리행
    spans = [[(c.colspan, c.rowspan) for c in row.cells] for row in table.rows]
    assert spans == [[(1, 1)] * 3, [(1, 2), (2, 1)], [(1, 1), (1, 1)]]


def test_docx_input_converts_end_to_end(word_file, tmp_path):
    out = tmp_path / "out.docx"
    result = convert_many([str(word_file)], out, "confluence")
    assert out.exists()
    texts = [p.text for p in DocxDocument(str(out)).paragraphs if p.text]
    assert texts[0] == "분기 보고"
    assert "□\t원문에 친 말머리" in texts


def test_docx_source_is_detected_by_extension(word_file):
    assert load_source(str(word_file)).format == "docx"


# ── 붙여넣은 글 ─────────────────────────────────────────────────────────


def test_pasted_lines_become_paragraphs_and_keep_typed_markers():
    doc = parse_markdown(text_to_markdown("1. 추진 배경\n□ 채용 현황\n- 입사 확정\n\n    들여쓴 줄"))
    texts = [plain(b.runs) for b in doc.blocks]
    assert texts == ["1. 추진 배경", "□ 채용 현황", "- 입사 확정", "들여쓴 줄"]
    assert all(isinstance(b, Paragraph) for b in doc.blocks)  # 목록 문법으로 먹히지 않음


def test_wrapped_line_stays_attached_to_the_marker_line_with_its_line_break():
    """원문에서 글쓴이가 엔터로 내려쓴 줄 — 말머리 줄 바로 아래에 이어지면 한 문장의 다음 줄이다(엔터 위치 보존)."""
    long = "핵심인력 선정 기준을 마련하여 인사 운영 체계를 개선하며,"
    md = text_to_markdown(f"1. 추진 배경\n□ {long}\n평가 결과를 공정하게 반영하고,\n성과 연계를 강화함\n"
                          f"- 짧은 항목\n다음 설명\n- 문장.\n마침표 뒤 줄\n*별표\n＊전각 별표\n(주석) 괄호\n"
                          f"2026. 9. 29\n날짜 다음 줄\n")
    texts = [plain(b.runs) for b in parse_markdown(md).blocks]
    assert texts == ["1. 추진 배경",
                     f"□ {long}\n평가 결과를 공정하게 반영하고,\n성과 연계를 강화함",   # 쉼표로 끝나는 줄 다음 줄도 이어진다
                     "- 짧은 항목\n다음 설명",                                     # 줄이 짧아도 이어진다
                     "- 문장.", "마침표 뒤 줄",                                    # 마침표로 끝난 줄 다음은 새 문단
                     "*별표", "＊전각 별표", "(주석) 괄호",                          # 주석 표시는 윗줄에 안 붙는다
                     "2026. 9. 29", "날짜 다음 줄"]                               # 날짜는 말머리 줄이 아니다


def test_tab_separated_lines_become_a_table():
    doc = parse_markdown(text_to_markdown("이름\t부서\n홍길동\t인사팀\n김철수\t재무팀\n"))
    (table,) = doc.blocks
    assert isinstance(table, Table) and len(table.rows) == 3 and table.header_rows == 1


def test_markdown_paste_is_left_alone():
    text = "# 제목\n\n- 항목\n- 항목2\n"
    assert looks_like_markdown(text) and text_to_markdown(text) == text


def test_pasted_markers_set_levels_without_headings(tmp_path):
    prof = load_profile("confluence")
    result = convert_many([load_text("1. 추진 배경\n□ 채용 현황\n- 입사 확정 3명")], None, prof)
    items = [(b.depth, b.marker) for b in result.document.blocks if isinstance(b, ListItem)]
    assert items == [(0, "1."), (1, "□"), (2, "-")]


def test_unstructured_memo_gets_markers_when_auto_markers_on():
    prof = load_profile("default")
    result = convert_many([load_text("2026. 9. 29\n첫 번째 내용입니다\n두 번째 내용입니다")], None, prof,
                          polish="none")
    blocks = result.document.blocks
    assert isinstance(blocks[0], Paragraph)  # 날짜 줄은 그대로
    assert [(type(b).__name__, b.depth) for b in blocks[1:]] == [("ListItem", 1), ("ListItem", 1)]


# ── 여러 입력 합치기 ────────────────────────────────────────────────────


def test_merged_inputs_get_section_titles_and_shifted_levels():
    a = load_text("□ 가 항목\n- 가 세부")
    a.title = "첫 페이지"
    b = load_text("□ 나 항목")
    b.title = "둘째 페이지"
    result = convert_many([a, b], None, "confluence", title="통합 보고")
    doc = result.document
    assert doc.title == "통합 보고"
    rows = [(b.depth, b.marker, plain(b.runs)) for b in doc.blocks if isinstance(b, ListItem)]
    assert rows == [(0, "", "첫 페이지"), (2, "□", "가 항목"), (3, "-", "가 세부"),
                    (0, "", "둘째 페이지"), (2, "□", "나 항목")]


def test_merge_without_section_titles_keeps_levels():
    a, b = load_text("□ 가"), load_text("□ 나")
    doc = convert_many([a, b], None, "confluence", section_titles=False).document
    assert [b.depth for b in doc.blocks if isinstance(b, ListItem)] == [0, 0]  # 같은 단계 (맨 바깥이라 0)


# ── Markdown / PDF 출력 ────────────────────────────────────────────────


def test_markdown_output_keeps_markers_bold_and_tables(tmp_path):
    prof = load_profile("confluence")
    src = load_text("1. 추진 배경\n□ 채용 현황\n※ 참고\n\n이름\t부서\n홍길동\t인사팀")
    src.title = "메모"
    result = convert_many([src], None, prof)
    text = render_markdown(result.document, prof)
    assert text.startswith("# 메모")
    assert "1\\. 추진 배경" in text  # "1. "이 Markdown 번호 목록으로 읽히지 않게
    assert "　□ 채용 현황" in text
    assert "| 이름 | 부서 |" in text
    path = write_markdown(result.document, prof, tmp_path / "x.md")
    assert path.read_text(encoding="utf-8") == text


def test_markdown_table_fills_merged_cells(word_file, tmp_path):
    prof = load_profile("confluence")
    doc = parse_docx(word_file, image_dir=tmp_path).document
    table = next(b for b in doc.blocks if isinstance(b, Table))
    from doc2report.render.markdown_writer import _grid

    assert _grid(table) == [["구분", "값", "비고"], ["세로병합", "가로병합", ""], ["", "x", "y"]]
    assert render_markdown(doc, prof)


def test_pdf_uses_libreoffice_when_word_is_absent(monkeypatch, tmp_path):
    from doc2report.render import pdf

    docx = tmp_path / "a.docx"
    docx.write_bytes(b"PK")
    calls = []

    def fake_run(cmd, **kwargs):
        calls.append(cmd)
        outdir = cmd[cmd.index("--outdir") + 1]
        (tmp_path / outdir / "a.pdf").write_bytes(b"%PDF")
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(pdf.platform, "system", lambda: "Linux")
    monkeypatch.setattr(pdf.shutil, "which", lambda name: "/usr/bin/soffice" if name == "soffice" else None)
    monkeypatch.setattr(pdf.subprocess, "run", fake_run)
    out = pdf.docx_to_pdf(docx, tmp_path / "결과.pdf")
    assert out.read_bytes() == b"%PDF" and calls[0][0] == "/usr/bin/soffice"


def test_pdf_reports_missing_converter(monkeypatch, tmp_path):
    from doc2report.render import pdf

    monkeypatch.setattr(pdf.platform, "system", lambda: "Linux")
    monkeypatch.setattr(pdf.shutil, "which", lambda name: None)
    with pytest.raises(RuntimeError, match="Word 또는 LibreOffice"):
        pdf.docx_to_pdf(tmp_path / "a.docx", tmp_path / "a.pdf")


def test_page_mode_puts_each_input_on_new_page_with_big_title(tmp_path):
    """입력마다 새 쪽: 쪽마다 그 입력 제목을 문서 제목 서식(큰 글씨·가운데·밑줄)으로, 날짜도 쪽마다
    제목 아래에(2026-09-29 사용자). 새 쪽 맨 위에 빈 줄이 생기지 않게 '쪽 나눔 앞'으로 나눈다."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    a, b = load_text("□ 가 항목"), load_text("□ 나 항목")
    a.title, b.title = "첫 보고", "둘째 보고"
    prof = load_profile("confluence")
    out = tmp_path / "p.docx"
    convert_many([a, b], out, prof, page_breaks=True, date="2026. 10. 1")
    paras = [p for p in DocxDocument(str(out)).paragraphs if p.text]
    texts = [p.text for p in paras]
    assert texts == ["첫 보고", "2026. 10. 1", "□\t가 항목", "둘째 보고", "2026. 10. 1", "□\t나 항목"]
    title_size = prof.font("title").size
    for p in (paras[0], paras[3]):
        run = p.runs[0]
        assert run.font.size == title_size and run.bold and run.underline
        assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paras[3].paragraph_format.page_break_before
    assert not any(p.text == "" and "w:br" in p._p.xml for p in DocxDocument(str(out)).paragraphs)
    assert paras[1].alignment == paras[4].alignment == WD_ALIGN_PARAGRAPH.RIGHT  # 날짜 서식
