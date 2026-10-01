"""정식보고서 서식(formal.yaml) — 공백으로 단계 구분, 왼쪽 정렬 표, 파란 주석.

목적이 다른 default·confluence 서식은 그대로여야 한다(2026-10-01 사용자).
"""

from __future__ import annotations

import pytest
from docx import Document as OpenDocx
from docx.shared import Emu, Pt, RGBColor

from doc2report.ir import Document, Heading, Paragraph, Run
from doc2report.pipeline import convert_document
from doc2report.profile import load_profile
from doc2report.web import options


def _doc() -> Document:
    return Document(blocks=[
        Heading(level=2, runs=[Run("추진 배경")]),
        Paragraph(runs=[Run("□ 핵심인력 선정")]),
        Paragraph(runs=[Run("- 세부 내용")]),
        Paragraph(runs=[Run("* 최근 평가 상위 이상, 영어회화 2급 이상")]),
        Paragraph(runs=[Run("□ 다음 항목")]),
        Paragraph(runs=[Run("※ 참고 사항")]),
    ])


@pytest.fixture()
def rendered(tmp_path):
    out = tmp_path / "formal.docx"
    convert_document(_doc(), out, load_profile("formal"), polish="none")
    return OpenDocx(str(out)).paragraphs


def test_formal_profile_values():
    prof = load_profile("formal")
    assert [(lv.lead_spaces, lv.marker_sep, lv.indent, lv.hanging) for lv in prof.numbering] == [
        (0, " ", 0, 0), (2, " ", 0, 0), (4, " ", 0, 0), (6, " ", 0, 0)]
    assert prof.tables.align == "left" and prof.tables.header_shading is None
    assert prof.text.annotation_markers == ["*"] and prof.text.note_lead_spaces == 4
    assert prof.font("annotation").color == "0000FF" and prof.font("annotation").size == Pt(10)


def test_other_profiles_are_unchanged():
    for name in ("default", "confluence"):
        prof = load_profile(name)
        assert [(lv.lead_spaces, lv.marker_sep) for lv in prof.numbering] == [(0, "\t")] * 4
        assert prof.numbering[1].indent == load_profile("default").numbering[1].indent > 0
        assert prof.tables.align == "right" and prof.tables.header_shading == "F2F2F2"
        assert prof.text.annotation_markers == [] and prof.text.note_lead_spaces is None
        assert not prof.has_font("annotation")


def test_levels_are_separated_by_spaces_not_indent(rendered):
    texts = [p.text for p in rendered]
    assert texts[:3] == ["1. 추진 배경", "  □ 핵심인력 선정", "    - 세부 내용"]
    for paragraph in rendered[:3]:
        assert not paragraph.paragraph_format.left_indent        # 들여쓰기 기능은 안 씀
        assert not paragraph.paragraph_format.first_line_indent  # 둘째 줄도 왼쪽 여백에서
    assert "\t" not in "".join(texts)


def test_annotation_is_blue_10pt_paragraph_not_a_list_item(rendered):
    note = next(p for p in rendered if "영어회화" in p.text)
    assert note.text == "    * 최근 평가 상위 이상, 영어회화 2급 이상"   # 앞 공백 4칸, "□"가 안 붙음
    run = next(r for r in note.runs if "영어회화" in r.text)
    assert run.font.size == Pt(10) and run.font.color.rgb == RGBColor(0x00, 0x00, 0xFF)
    assert run.font.name == "바탕체" or run._r.rPr.rFonts.get(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia") == "바탕체"


def test_annotation_keeps_level_gaps_of_the_line_above(rendered):
    prof = load_profile("formal")
    by_text = {p.text.strip(): p for p in rendered}
    above, note = by_text["- 세부 내용"], by_text["* 최근 평가 상위 이상, 영어회화 2급 이상"]
    assert above.paragraph_format.space_after == Emu(0)                       # 주석이 윗줄에 붙는다
    assert note.paragraph_format.space_after == prof.numbering[2].space_after_level_change  # - → □ 간격


def test_note_mark_uses_spaces_and_body_size(rendered):
    note = next(p for p in rendered if "참고 사항" in p.text)
    assert note.text == "    ※ 참고 사항"
    assert all(r.font.size in (None, Pt(14)) for r in note.runs)               # 본문보다 작게 안 함


def test_web_preset_keeps_format_text_values():
    docs = [_doc()]
    for name, expected in (("formal", ["*"]), ("confluence", [])):
        prof = options.build_profile({"preset": name}, docs, ["text"], llm_ready=False)[0]
        assert prof.text.annotation_markers == expected
    prof = options.build_profile({"preset": "formal", "mode": "manual"}, docs, ["text"], llm_ready=False)[0]
    assert prof.text.note_lead_spaces == 4 and prof.numbering[1].lead_spaces == 2
    assert "formal" in options.presets()


_TABLE_MD = """## 현황

| 구분 | 내용 |
|---|---|
| 현황 | 양호함 |
"""


def _table_xml(tmp_path, profile: str):
    src = tmp_path / "t.md"
    src.write_text(_TABLE_MD, encoding="utf-8")
    out = tmp_path / f"{profile}.docx"
    from doc2report.pipeline import convert

    convert(str(src), out, profile, polish="none")
    return OpenDocx(str(out)).tables[0]._tbl.xml


def test_formal_table_is_left_aligned_without_shading(tmp_path):
    xml = _table_xml(tmp_path, "formal")
    assert 'w:jc w:val="left"' in xml and "w:shd" not in xml


def test_default_table_keeps_right_alignment_and_header_shading(tmp_path):
    xml = _table_xml(tmp_path, "default")
    assert 'w:jc w:val="right"' in xml and "F2F2F2" in xml
