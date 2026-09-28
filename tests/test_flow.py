"""지면 여유 판단 — "공간이 충분하면 넉넉하게" 규칙의 근거.

기본 프로파일은 단계 전환 간격을 6pt로 고정했다(사용자 요청, 2026-09-28).
'여유 있으면 늘리기'는 space_after_level_change_max를 준 프로파일에서만 동작하므로,
그 규칙 자체를 검사하는 테스트는 그 값을 넣은 프로파일로 돌린다.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.oxml.ns import qn

from doc2report.ir import Document, ListItem, Run
from doc2report.layout.flow import plan_flow
from doc2report.pipeline import convert
from doc2report.profile import load_profile
from doc2report.units import parse_length

FIXTURES = Path(__file__).parent / "fixtures"
ITEM_MARKERS = ("1.\t", "2.\t", "3.\t", "4.\t", "□\t", "-\t")


@pytest.fixture
def profile():
    return load_profile("default")


@pytest.fixture
def stretchy(profile):
    """단계 전환 간격을 여유 있으면 12pt까지 늘리는 프로파일."""
    wider = parse_length("12pt")
    levels = [level.model_copy(update={"space_after_level_change_max": wider})
              for level in profile.numbering]
    return profile.model_copy(update={"numbering": levels})


def _item_spacings(path: Path) -> set[int]:
    docx = DocxDocument(str(path))
    return {int(p.paragraph_format.space_after or 0) for p in docx.paragraphs
            if p.text.startswith(ITEM_MARKERS)}


def test_level_change_gap_is_6pt_by_default(tmp_path, profile):
    """1.→□, □→-/표 자리의 단락 뒤 간격은 기본 6pt — 지면이 남아도 늘리지 않는다."""
    out = tmp_path / "sample.docx"
    convert(str(FIXTURES / "sample_report.md"), out, profile)
    six = parse_length("6pt")
    assert _item_spacings(out) <= {0, six}
    assert six in _item_spacings(out)


def test_short_document_has_room_to_breathe(stretchy):
    doc = Document(blocks=[ListItem(depth=0, runs=[Run("짧은 항목")]),
                           ListItem(depth=1, runs=[Run("하위 항목")])],
                   title="짧은 보고서")
    plan = plan_flow(doc, stretchy, {})
    assert plan.pages == 1
    assert plan.relaxed is True
    assert plan.notes


def test_nearly_full_page_stays_tight(stretchy):
    """쪽을 거의 채운 문서는 간격을 늘리지 않는다 (늘리면 쪽이 넘어가므로)."""
    usable = stretchy.page.usable_height
    blocks: list = []
    plan = plan_flow(Document(blocks=blocks), stretchy, {})
    while plan.estimated_height < usable * 0.92:
        depth = len(blocks) % 3
        blocks.append(ListItem(depth=depth, runs=[Run(f"항목 {len(blocks)}")]))
        plan = plan_flow(Document(blocks=blocks), stretchy, {})
    assert plan.pages == 1
    assert plan.relaxed is False


def test_estimate_is_close_to_reality(tmp_path, profile):
    """어림한 쪽 수가 실제 렌더 결과와 맞아야 판단을 믿을 수 있다."""
    out = tmp_path / "company.docx"
    result = convert(str(FIXTURES / "company_format.md"), out, profile)
    assert result.flow.pages == 1  # Word 렌더 결과도 1쪽 (tools/score_corpus.py로 확인)
    assert result.flow.estimated_height < profile.page.usable_height


def test_relaxed_document_gets_taller_rows(tmp_path, profile):
    out = tmp_path / "wide.docx"
    result = convert(str(FIXTURES / "wide_table.md"), out, profile)
    assert result.flow.relaxed is True

    table = DocxDocument(str(out)).tables[0]
    height = table.rows[0]._tr.find(qn("w:trPr")).find(qn("w:trHeight"))
    assert int(height.get(qn("w:val"))) == pytest.approx(
        profile.tables.row_height_relaxed / 635, abs=1)


def test_relaxed_document_gets_wider_gaps_when_profile_allows(tmp_path, stretchy):
    out = tmp_path / "wide.docx"
    convert(str(FIXTURES / "wide_table.md"), out, stretchy)
    assert parse_length("12pt") in _item_spacings(out)
