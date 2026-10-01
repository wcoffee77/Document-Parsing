"""정식보고서 줄 맞춤·주석 정규화·굵기·txt 제목 인식 (2026-10-01 사용자 실측 반영)."""

from __future__ import annotations

import pytest
from docx import Document as OpenDocx
from docx.shared import Pt, RGBColor

from doc2report.ir import Document, Heading, Paragraph, Run
from doc2report.layout.lines import fit_text
from doc2report.parsers.plaintext import text_to_markdown
from doc2report.pipeline import convert_document
from doc2report.profile import load_profile
from doc2report.render import docx_writer
from doc2report.transform.structure import normalize_annotations

W = {"narrow": 5.0, "wide": 10.0}  # 한글 10, 반각(공백·숫자) 5 — 바탕체처럼


def widths(text: str) -> list[float]:
    return [W["narrow"] if c == " " or c.isascii() else W["wide"] for c in text]


# ── 줄 나눔 알고리즘 ─────────────────────────────────────────────────

def test_short_text_stays_one_line():
    text = "가나다 라마바"
    lines = fit_text(text, widths(text), first_room=200, cont_room=200)
    assert [(l.start, l.end, l.steps) for l in lines] == [(0, len(text), 0)]


def test_slight_overflow_is_condensed_not_wrapped():
    text = "가나다 라마바"                                   # 폭 = 6*10 + 5 = 65
    lines = fit_text(text, widths(text), first_room=62, cont_room=62, max_condense=1.0, step=0.25)
    assert len(lines) == 1 and lines[0].steps == 2         # 넘침 3 / 글자 7개 = 0.43 → 0.25 단위로 올림 = 2단계(0.5)
    lines = fit_text(text, widths(text), first_room=40, cont_room=40, max_condense=1.0, step=0.25)
    assert len(lines) == 2                                  # 너무 많이 넘치면 줄을 나눈다


def test_wraps_at_word_boundary_with_continuation_room():
    text = "가나다 라마바 사아자 차카타"
    lines = fit_text(text, widths(text), first_room=75, cont_room=95)
    assert [text[l.start:l.end] for l in lines] == ["가나다 라마바", "사아자 차카타"]


def test_overlong_word_is_split_by_characters():
    text = "가나다라마바사아"
    lines = fit_text(text, widths(text), first_room=35, cont_room=35)
    assert [text[l.start:l.end] for l in lines] == ["가나다", "라마바", "사아"]


def test_margin_keeps_the_line_shorter():
    text = "가나다 라마바"
    assert len(fit_text(text, widths(text), first_room=65, cont_room=65)) == 1
    assert len(fit_text(text, widths(text), first_room=65, cont_room=65, margin=0.1)) == 2


# ── 주석 정규화 ──────────────────────────────────────────────────────

def test_annotation_marker_is_unified_and_text_rules_skip_it():
    doc = Document(blocks=[Paragraph(runs=[Run("(주석) 최근 평가 상위 이상")]),
                           Paragraph(runs=[Run("* 이미 표준 표시")])])
    out, changes = normalize_annotations(doc, ["*", "(주석)"], "*")
    assert [b.runs[0].text + "".join(r.text for r in b.runs[1:]) for b in out.blocks] == [
        "* 최근 평가 상위 이상", "* 이미 표준 표시"]
    assert [c.rule for c in changes] == ["주석 표시 통일"]


# ── 렌더: 굵기·주석·줄 맞춤 ───────────────────────────────────────────

class _FakeMeasurer:
    """바탕체처럼 한글 1em, 반각 0.5em (글꼴 파일이 없는 환경용)."""
    font_available = True

    def __init__(self, spec, bold):
        self.size = spec.size

    def char_width(self, ch):
        return self.size * (0.5 if ch == " " or ch.isascii() else 1.0)

    def width(self, text):
        return sum(self.char_width(c) for c in text)


@pytest.fixture()
def fake_fonts(monkeypatch):
    monkeypatch.setattr(docx_writer.DocxRenderer, "_measurer",
                        lambda self, spec, bold: _FakeMeasurer(spec, bold))


def _render(tmp_path, blocks):
    out = tmp_path / "o.docx"
    convert_document(Document(blocks=blocks), out, load_profile("formal"), polish="none")
    return OpenDocx(str(out)).paragraphs


def test_original_markers_are_bold_for_formal(tmp_path):
    paragraphs = _render(tmp_path, [Heading(level=2, runs=[Run("추진 배경")]),
                                    Paragraph(runs=[Run("□ 핵심인력 선정")]),
                                    Paragraph(runs=[Run("- 세부 내용")])])
    bold = {p.text.strip(): all(r.bold for r in p.runs) for p in paragraphs}
    assert bold["1. 추진 배경"] and bold["□ 핵심인력 선정"] and not bold["- 세부 내용"]


def test_other_profiles_keep_original_marker_weight(tmp_path):
    out = tmp_path / "d.docx"
    convert_document(Document(blocks=[Heading(level=2, runs=[Run("추진")]),
                                      Paragraph(runs=[Run("□ 항목")])]),
                     out, load_profile("default"), polish="none")
    item = next(p for p in OpenDocx(str(out)).paragraphs if "항목" in p.text)
    assert not any(r.bold for r in item.runs)


def test_paren_annotation_becomes_blue_annotation(tmp_path):
    paragraphs = _render(tmp_path, [Heading(level=2, runs=[Run("추진 배경")]),
                                    Paragraph(runs=[Run("□ 핵심인력")]),
                                    Paragraph(runs=[Run("(주석) 최근 평가 상위 이상")])])
    note = next(p for p in paragraphs if "최근 평가" in p.text)
    assert note.text == "    * 최근 평가 상위 이상"
    run = next(r for r in note.runs if "최근" in r.text)
    assert run.font.size == Pt(10) and run.font.color.rgb == RGBColor(0, 0, 255)


def test_long_item_is_split_and_aligned_under_text(tmp_path, fake_fonts):
    long = " ".join(["가나다라마바사아"] * 12)                      # 한 줄(약 170mm)을 넘는 문장
    paragraphs = _render(tmp_path, [Heading(level=2, runs=[Run("추진 배경")]),
                                    Paragraph(runs=[Run("□ " + long)])])
    items = [p for p in paragraphs if p.text.strip() and not p.text.startswith("1.")]
    assert len(items) >= 2
    first, second = items[0], items[1]
    assert first.text.startswith("  □ ")
    # 접두 "  □ " = 반각 2 + 전각 □(반각 2) + 반각 1 = 5칸 → 둘째 줄은 공백 5칸으로 글자 시작에 맞춘다
    assert second.text.startswith(" " * 5) and not second.text.startswith(" " * 6)
    assert first.paragraph_format.keep_with_next and not items[-1].paragraph_format.keep_with_next


def test_slight_overflow_condenses_instead_of_wrapping(tmp_path, fake_fonts):
    prof = load_profile("formal")
    usable = prof.page.usable_width
    per = prof.font("body").size                                    # 전각 1글자 = 1em
    count = int(usable / per) - 3                                   # 접두("  □ ", 공백 포함) 폭을 빼면 딱 1글자쯤 넘침
    body = "가" * (count + 1)
    paragraphs = _render(tmp_path, [Heading(level=2, runs=[Run("추진 배경")]),
                                    Paragraph(runs=[Run("□ " + body)])])
    items = [p for p in paragraphs if "가" in p.text]
    assert len(items) == 1                                          # 줄을 안 나눴다
    spacing = [r._r.rPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing") for r in items[0].runs]
    values = [int(x.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"))
              for x in spacing if x is not None]
    assert values and all(-10 <= v <= -2 for v in values)           # 0.1~0.5pt = -2~-10 (1/20pt)


# ── 서식 없는 txt: 제목·날짜 ──────────────────────────────────────────

def test_txt_title_and_date_are_recognised():
    md = text_to_markdown("인재 운영 현황 보고\n\n2026. 10. 1\n\n1. 추진 배경\n")
    assert md.splitlines()[0] == "# 인재 운영 현황 보고"
    assert not text_to_markdown("1. 추진 배경\n2026. 10. 1\n").startswith("# ")   # 말머리 줄은 제목이 아님


def test_star_right_after_a_table_is_still_a_blue_annotation(tmp_path):
    from doc2report.pipeline import convert

    src = tmp_path / "s.txt"
    src.write_text("1. 현황\n\n구분\t현황\nA\t양호\n* 표 아래 설명\n※ 표 아래 참고\n", encoding="utf-8")
    out = tmp_path / "s.docx"
    convert(str(src), out, "formal", polish="none")
    paragraphs = {p.text.strip(): p for p in OpenDocx(str(out)).paragraphs}
    star = paragraphs["* 표 아래 설명"]
    run = star.runs[-1]
    assert run.font.size == Pt(10) and run.font.color.rgb == RGBColor(0, 0, 255)
    note = paragraphs["※ 표 아래 참고"]
    assert note.runs[-1].font.size in (None, Pt(14)) and not (
        note.runs[-1].font.color and note.runs[-1].font.color.type and
        note.runs[-1].font.color.rgb == RGBColor(0, 0, 255))
