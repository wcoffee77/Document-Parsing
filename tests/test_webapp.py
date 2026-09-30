"""웹 화면 — 옵션→프로파일, 자동 판단, 작업 실행·파일 이름, HTTP 엔드포인트."""

from __future__ import annotations

import json
import threading
import urllib.error
import urllib.request

import pytest
from docx import Document as DocxDocument

from doc2report.web import options as opts
from doc2report.web.jobs import JobRunner, history, safe_name
from doc2report.web.server import create_server


# ── 옵션 ────────────────────────────────────────────────────────────────


def test_presets_are_profiles_with_preset_order():
    assert opts.presets()[:2] == ["default", "confluence"]
    report, confluence = opts.profile_info("default"), opts.profile_info("confluence")
    assert (report["label"], confluence["label"]) == ("보고서", "Confluence 변환")


def test_confluence_preset_matches_user_spec():
    """2026-09-29 사용자(두 번째 확정): 여백 모두 2cm, 맑은 고딕, 제목 18pt, 본문 12pt·장평 95%·
    줄간격 1.0, 표 10pt·장평 80%."""
    fmt = opts.profile_info("confluence")["format"]
    assert (fmt["margin_top"], fmt["margin_bottom"], fmt["margin_left"], fmt["margin_right"]) == \
        ("2", "2", "2", "2")
    assert (fmt["title_size"], fmt["size"], fmt["font"], fmt["line_spacing"]) == \
        ("18pt", "12pt", "맑은 고딕", "1.0")
    assert (fmt["body_scale"], fmt["table_size"], fmt["table_scale"]) == ("95%", "10pt", "80%")
    prof = opts.load_profile("confluence")
    assert prof.font("table").east_asia == "맑은 고딕"  # 제목·표도 body 글꼴을 물려받음


def test_custom_preset_sets_body_and_table_char_scale():
    prof = opts.format_profile({"preset": "custom", "custom": {
        "base": "default", "body_scale": "90", "table_scale": "85%"}})
    assert prof.font("body").char_scale == pytest.approx(0.9)  # % 없이 적어도 90%
    assert [scale for _, scale in prof.table_steps()] == [pytest.approx(0.85)] * len(prof.table_steps())


def test_custom_preset_overrides_only_changed_values():
    prof = opts.format_profile({"preset": "custom", "custom": {
        "base": "confluence", "size": "13pt", "margin_top": "1.2", "font": "맑은 고딕"}})
    base = opts.load_profile("confluence")
    assert prof.font("body").size == 13 * 12700
    assert prof.page.margin.top == 12 * 36000 and prof.page.margin.left == base.page.margin.left
    assert prof.font("title").size == base.font("title").size
    with pytest.raises(ValueError):
        opts.format_profile({"preset": "../../etc"})
    with pytest.raises(ValueError):
        opts.format_profile({"preset": "custom", "custom": {"base": "nope"}})


def test_manual_rules_are_put_into_the_chosen_preset():
    prof, polish, llm, _ = opts.build_profile({
        "mode": "manual", "preset": "confluence", "rules_base": "default",
        "polish": "rules", "llm": True,
        "text": {"auto_markers": True, "endings": False, "unknown": True},
        "tables": {"allow_landscape": True, "align": "center"},
    }, [], [], llm_ready=False)
    assert (polish, llm) == ("rules", True)
    assert prof.text.auto_markers
    assert not prof.text.gaechosik and not prof.text.noun_ending  # 한 체크박스가 두 규칙을 함께
    assert prof.text.level_bold  # 화면에 없는 규칙은 "규칙 기본값"(보고서) 값
    assert prof.tables.allow_landscape and prof.tables.align == "center"
    assert prof.font("body").east_asia == "맑은 고딕"  # 서식은 preset 그대로
    assert opts.profile_info("confluence")["text"]["auto_markers"] is False  # 프로파일 파일은 그대로
    with pytest.raises(ValueError):
        opts.build_profile({"mode": "manual", "polish": "magic"}, [], [], llm_ready=False)


def test_every_toggle_is_a_real_profile_field():
    fields = set(opts.load_profile("default").text.model_dump())
    for _, _, _, targets in opts.MARKER_TOGGLES + opts.POLISH_TOGGLES + opts.TABLE_TOGGLES:
        assert set(targets) <= fields


def test_rules_that_are_not_on_screen_come_from_rules_base():
    """제목 접기·0cm·단계 굵게·표 제목은 화면에서 뺐다(2026-09-29 사용자) — 규칙 기본값을 따른다."""
    keys = {k for k, _, _, _ in opts.MARKER_TOGGLES + opts.POLISH_TOGGLES + opts.TABLE_TOGGLES}
    assert not keys & {"headings_as_levels", "normalize_levels", "level_bold", "table_captions",
                       "gaechosik", "noun_ending"}
    conf, _, _, _ = opts.build_profile({"mode": "manual", "rules_base": "confluence"}, [], [],
                                       llm_ready=False)
    assert conf.text.headings_as_levels and conf.text.table_captions and not conf.text.level_bold


def test_polish_none_can_still_use_llm():
    _, polish, llm, _ = opts.build_profile({"mode": "manual", "polish": "none", "llm": True},
                                           [], [], llm_ready=True)
    assert (polish, llm) == ("none", True)
    _, polish, llm, _ = opts.build_profile({"mode": "manual", "polish": "llm"}, [], [],
                                           llm_ready=True)
    assert (polish, llm) == ("rules", True)  # 예전 값(규칙+LLM)도 받아 준다


def test_auto_decides_structured_confluence_vs_unstructured_memo():
    from doc2report.pipeline import load_document
    from doc2report.sources import load_text

    structured, _ = load_document(load_text("1. 배경\n□ 현황\n- 세부"))
    memo, _ = load_document(load_text("오늘 회의에서 채용 이야기를 했습니다\n다음 주에 다시 봅니다"))

    keep = opts.auto_decide([structured], ["text"], allow_llm=True, llm_ready=True)
    assert keep.polish == "none" and keep.profile.text.auto_markers is False

    tidy = opts.auto_decide([memo], ["text"], allow_llm=True, llm_ready=True)
    assert (tidy.polish, tidy.llm) == ("rules", True) and tidy.profile.text.auto_markers is True
    assert opts.auto_decide([memo], ["text"], allow_llm=True, llm_ready=False).llm is False
    assert opts.auto_decide([memo], ["text"], allow_llm=False, llm_ready=True).llm is False

    confluence = opts.auto_decide([memo], ["confluence"], allow_llm=True, llm_ready=True)
    assert confluence.polish == "none"  # Confluence는 파이썬 규칙으로 다듬지 않음(사용자 규칙)
    assert confluence.llm is True       # LLM 맞춤법·어조는 사용자가 켜면 적용
    two = opts.auto_decide([memo, memo], ["text", "text"], allow_llm=False, llm_ready=False)
    assert two.summary["heavy"]
    assert two.profile.font("body").size == opts.load_profile("default").font("body").size  # 서식은 그대로
    assert any("Confluence 변환" in r for r in two.reasons)  # 권하기만


def test_auto_rules_keep_the_chosen_preset_format():
    from doc2report.pipeline import load_document
    from doc2report.sources import load_text

    memo, _ = load_document(load_text("회의 메모입니다"))
    prof, polish, _, decision = opts.build_profile({"preset": "confluence"}, [memo], ["text"],
                                                   llm_ready=False)
    assert prof.font("body").east_asia == "맑은 고딕" and prof.text.auto_markers is True
    assert decision is not None and polish == "rules"


def test_date_option_is_formatted_with_profile():
    prof = opts.load_profile("default")
    assert opts.date_text("", prof) is None
    assert opts.date_text("today", prof) == "today"
    assert opts.date_text("2026-10-01", prof) == "2026. 10. 1"


# ── 작업·파일 이름 ──────────────────────────────────────────────────────


def test_safe_name_strips_windows_illegal_characters():
    assert safe_name('3분기: 보고/정리 "초안"?') == "3분기_보고_정리_초안"
    assert safe_name("   ") == "보고서"


def test_job_writes_all_formats_with_unique_names(tmp_path):
    runner = JobRunner(tmp_path / "out", tmp_path / "up")
    payload = {"inputs": [{"type": "text", "text": "1. 배경\n□ 현황", "title": "메모"}],
               "options": {"mode": "manual", "preset": "confluence", "formats": ["docx", "md"],
                           "title": "주간: 보고", "date": "today"}}
    first = runner.run_sync(payload)
    second = runner.run_sync(payload)
    assert first.state == "done", first.error
    names = [f["name"] for f in first.result["files"]]
    stem = first.result["stem"]
    assert stem.endswith("_주간_보고")
    assert names == [f"{stem}.docx", f"{stem}.md", f"{stem}_변경내역.md"]
    assert second.result["stem"] == stem + "_2"  # 같은 이름이 있으면 번호를 붙인다
    groups = history(tmp_path / "out")
    assert {g["stem"] for g in groups} == {stem, stem + "_2"}
    assert all({f["kind"] for f in g["files"]} == {"docx", "md", "report"} for g in groups)


def test_job_reports_bad_input(tmp_path):
    runner = JobRunner(tmp_path / "out", tmp_path / "up")
    job = runner.run_sync({"inputs": [], "options": {}})
    assert job.state == "error" and "입력이 없습니다" in job.error


def test_pdf_failure_keeps_word_output(tmp_path, monkeypatch):
    import doc2report.render.pdf as pdf

    def boom(*a, **k):
        raise RuntimeError("PDF 변환기 없음")

    monkeypatch.setattr(pdf, "docx_to_pdf", boom)
    runner = JobRunner(tmp_path / "out", tmp_path / "up")
    job = runner.run_sync({"inputs": [{"type": "text", "text": "□ 가"}],
                           "options": {"formats": ["docx", "pdf"]}})
    assert job.state == "done"
    assert [f["kind"] for f in job.result["files"]] == ["docx", "report"]
    assert any("PDF 만들기 실패" in n for n in job.result["notes"])


# ── HTTP ────────────────────────────────────────────────────────────────


@pytest.fixture
def server(tmp_path):
    srv, app = create_server("127.0.0.1", 0, tmp_path / "out")
    thread = threading.Thread(target=srv.serve_forever, daemon=True)
    thread.start()
    yield f"http://127.0.0.1:{srv.server_address[1]}", app
    srv.shutdown()
    srv.server_close()


def _call(url, body=None, headers=None, raw=None):
    data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
    req = urllib.request.Request(url, data=data, headers=headers or {})
    try:
        with urllib.request.urlopen(req, timeout=30) as res:
            return res.status, res.headers, res.read()
    except urllib.error.HTTPError as err:
        return err.code, err.headers, err.read()


def _post(base, path, body=None, raw=None, headers=None):
    head = {"X-Doc2Report": "1", "Content-Type": "application/json", **(headers or {})}
    status, _, data = _call(base + path, body, head, raw)
    return status, json.loads(data)


def test_http_pages_and_status(server):
    base, app = server
    status, headers, body = _call(base + "/")
    assert status == 200 and "doc2report" in body.decode()
    status, headers, _ = _call(base + "/static/app.css")
    assert headers["Content-Type"].startswith("text/css")
    status, _, body = _call(base + "/api/status")
    data = json.loads(body)
    assert data["output_dir"] == str(app.output_dir)
    assert "token" not in json.dumps(data).lower()
    profiles = json.loads(_call(base + "/api/profiles")[2])
    assert {"default", "confluence"} <= {p["name"] for p in profiles["profiles"]}


def test_http_post_requires_custom_header(server):
    base, _ = server
    status, _, _ = _call(base + "/api/convert", {"inputs": []}, {"Content-Type": "application/json"})
    assert status == 403


def test_http_upload_convert_and_download(server, tmp_path):
    base, app = server
    word = tmp_path / "원본.docx"
    d = DocxDocument()
    d.add_paragraph("□ 워드 항목")
    d.save(str(word))

    status, up = _post(base, "/api/upload", raw=word.read_bytes(),
                       headers={"X-Filename": "%EC%9B%90%EB%B3%B8.docx",
                                "Content-Type": "application/octet-stream"})
    assert status == 200 and up["name"] == "원본.docx"
    status, bad = _post(base, "/api/upload", raw=b"not a zip", headers={"X-Filename": "x.docx"})
    assert status == 400 and not bad["ok"]

    inputs = [{"type": "docx", "upload_id": up["upload_id"], "name": "원본.docx"},
              {"type": "text", "text": "- 붙여넣은 항목", "title": "메모"}]
    status, started = _post(base, "/api/convert", {"inputs": inputs,
                                                   "options": {"mode": "auto", "formats": ["md"]}})
    job = app.jobs.jobs[started["job"]]
    for _ in range(200):
        if job.state != "running":
            break
        threading.Event().wait(0.05)
    data = json.loads(_call(base + f"/api/jobs/{job.id}")[2])
    assert data["state"] == "done", data
    assert data["result"]["decision"]["reasons"]
    names = [f["name"] for f in data["result"]["files"]]
    status, headers, body = _call(base + "/files/" + urllib.request.quote(names[0]))
    assert status == 200 and body[:2] == b"PK" and "attachment" in headers["Content-Disposition"]
    md = next(n for n in names if n.endswith(".md") and "변경내역" not in n)
    status, headers, body = _call(base + "/files/" + urllib.request.quote(md))
    assert "inline" in headers["Content-Disposition"] and "워드 항목" in body.decode("utf-8")
    items = json.loads(_call(base + "/api/history")[2])["items"]
    assert items and items[0]["stem"] == data["result"]["stem"]


def test_http_files_cannot_escape_output_folder(server, tmp_path):
    base, _ = server
    (tmp_path / "secret.txt").write_text("x")
    for path in ("/files/..%2Fsecret.txt", "/files/../secret.txt", "/files/%2Fetc%2Fpasswd"):
        assert _call(base + path)[0] == 404
    status, data = _post(base, "/api/open", {"name": "../secret.txt"})
    assert status == 400 and not data["ok"]


# ── 예전 서버가 남아 있는 문제 (2026-09-29 사용자 PC) ───────────────────


def test_status_reports_api_version_and_stale_code(server):
    from doc2report.web import server as srv

    base, app = server
    data = json.loads(_call(base + "/api/status")[2])
    assert data["version"] == srv.API_VERSION and data["stale"] is False
    app.fingerprint = "예전 코드"  # 서버를 켠 뒤 git pull로 코드가 바뀐 상황
    assert json.loads(_call(base + "/api/status")[2])["stale"] is True


def test_app_js_expects_the_same_api_version():
    from pathlib import Path

    from doc2report.web import server as srv

    js = (Path(srv.STATIC) / "app.js").read_text(encoding="utf-8")
    assert f"const API_VERSION = {srv.API_VERSION};" in js


def test_new_server_takes_over_port_from_running_one(tmp_path):
    from doc2report.web import server as srv

    first, _ = srv.create_server("127.0.0.1", 0, tmp_path / "a")
    port = first.server_address[1]
    thread = threading.Thread(target=first.serve_forever, daemon=True)
    thread.start()
    with pytest.raises(OSError):
        srv.create_server("127.0.0.1", port, tmp_path / "b")  # 포트가 쓰이고 있으면 확실히 실패
    assert srv._take_over(port) == "stopped"
    thread.join(timeout=5)
    assert not thread.is_alive()
    first.server_close()
    second, _ = srv.create_server("127.0.0.1", port, tmp_path / "b")  # 이제 잡힌다
    second.server_close()


def test_broken_profile_file_does_not_break_the_screen(tmp_path, monkeypatch):
    import shutil

    from doc2report import profile as profile_mod

    for name in ("default.yaml", "confluence.yaml"):
        shutil.copy(profile_mod.PROFILE_DIR / name, tmp_path / name)
    (tmp_path / "망가진.yaml").write_text("fonts: [이상한 값", encoding="utf-8")
    monkeypatch.setattr(profile_mod, "PROFILE_DIR", tmp_path)
    monkeypatch.setattr(opts, "PROFILE_DIR", tmp_path)
    assert opts.profile_names() == ["confluence", "default"]
    assert [e.split(":")[0] for e in opts.profile_errors()] == ["망가진.yaml"]


def test_confluence_and_word_inputs_keep_original_bold_even_with_report_rules():
    """규칙 기본값을 '보고서'로 골라도 Confluence·Word 입력이면 1.·□ 문장 전체 굵게를 끈다
    (2026-09-30 사용자: 원문에서 굵은 글씨만 굵게)."""
    from doc2report.ir import Document
    from doc2report.web.options import build_profile

    opts = {"mode": "manual", "preset": "confluence", "rules_base": "default", "polish": "none", "text": {}}
    for kinds, expected in ((["confluence"], False), (["docx"], False), (["text"], True)):
        prof, *_ = build_profile(opts, [Document(blocks=[])], kinds, llm_ready=False)
        assert prof.text.level_bold is expected
