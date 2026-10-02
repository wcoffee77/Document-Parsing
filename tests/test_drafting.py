"""줄글 → 보고서 구조(drafting.py): LLM은 배치만, 문장은 원문 그대로."""

import json

import pytest

from doc2report.drafting import draft, parse_structure, sanitize, split_sentences

TEXT = """충원 현황 정리

9월 말 기준으로 40명 중 27명이 입사를 확정했습니다. 달성률은 67.5%입니다. 설계 직군은 14명 중 8명만 확보했습니다.
헤드헌팅 수수료를 30%로 올렸습니다. 최종 비용은 달라질 수 있어 확인이 필요합니다.
"""


def _sentences():
    return split_sentences(TEXT)


def _ok(_system, _user):
    return json.dumps({"title": "충원 현황", "sections": [
        {"heading": "현황", "groups": [{"summary": "경력 채용 27명 입사 확정", "ids": [1, 2], "notes": []},
                                     {"summary": "", "ids": [3], "notes": []}]},
        {"heading": "대응", "groups": [{"summary": "수수료 30% 인상", "ids": [4], "notes": [5]}]}]}, ensure_ascii=False)


def test_split_sentences_title_and_dates():
    title, sentences = _sentences()
    assert title == "충원 현황 정리" and len(sentences) == 5
    assert "67.5%입니다." in sentences[1]                        # 소수점·퍼센트를 문장 끝으로 오인하지 않는다


def test_draft_keeps_original_sentences_and_places_them():
    result = draft(TEXT, ask=_ok)
    assert result.used_llm and not [n for n in result.notes if "실패" in n]
    lines = result.text.splitlines()
    assert "1. 현황" in lines and "2. 대응" in lines
    assert "□ 경력 채용 27명 입사 확정" in lines                  # 요지는 새로 쓴 것
    assert "- 9월 말 기준으로 40명 중 27명이 입사를 확정했습니다." in lines    # 세부는 원문 그대로
    assert "□ 설계 직군은 14명 중 8명만 확보했습니다." in lines    # 문장 하나뿐인 묶음은 문장이 □
    assert any(line.startswith("* 최종 비용은") for line in lines)  # 단서는 주석으로


def test_invented_number_in_summary_is_dropped():
    def bad(_s, _u):
        data = json.loads(_ok(_s, _u))
        data["sections"][0]["groups"][0]["summary"] = "경력 채용 35명 입사 확정"     # 원문에 없는 35
        return json.dumps(data, ensure_ascii=False)

    result = draft(TEXT, ask=bad)
    assert any("원문에 없는 숫자" in n for n in result.notes)
    assert "35명" not in result.text and "□ 9월 말 기준으로 40명" in result.text      # 요지를 버리고 첫 문장이 □


def test_missing_sentence_triggers_retry_then_fallback():
    calls = []

    def lossy(_s, user):
        calls.append(user)
        return json.dumps({"sections": [{"heading": "현황", "groups": [{"summary": "", "ids": [1, 2], "notes": []}]}]})

    result = draft(TEXT, ask=lossy)
    assert len(calls) == 2 and "이전 응답의 문제" in calls[1]
    assert not result.used_llm and "규칙 기본 구조" in " ".join(result.notes)
    for sentence in _sentences()[1]:
        assert sentence in result.text                            # 문장이 하나도 사라지지 않는다


def test_llm_error_falls_back_without_crashing():
    def boom(_s, _u):
        raise RuntimeError("연결 실패")

    result = draft(TEXT, ask=boom)
    assert not result.used_llm and any("연결 실패" in n for n in result.notes)


def test_order_must_follow_original():
    sentences = _sentences()[1]
    raw = json.dumps({"sections": [{"heading": "a", "groups": [{"summary": "", "ids": [2], "notes": []},
                                                                {"summary": "", "ids": [1, 3, 4, 5], "notes": []}]}]})
    with pytest.raises(ValueError, match="순서"):
        parse_structure(raw, sentences)


def test_draft_text_goes_through_formal_converter(tmp_path):
    from doc2report.pipeline import convert_many

    result = draft(TEXT, ask=_ok)
    src = tmp_path / "s.txt"
    src.write_text(result.text, encoding="utf-8")
    out = convert_many([str(src)], tmp_path / "o.docx", "formal", title=result.structure.title, section_titles=False)
    from docx import Document

    body = "\n".join(p.text for p in Document(str(tmp_path / "o.docx")).paragraphs)
    assert "경력 채용 27명 입사 확정" in body and "수수료를 30%로 올렸습니다" in body


def test_title_is_not_duplicated_in_converted_document(tmp_path):
    from docx import Document

    from doc2report.pipeline import convert_many

    result = draft(TEXT, ask=_ok)
    assert result.structure.title == "충원 현황 정리"           # 원문 제목이 있으면 원문 그대로
    src = tmp_path / "s.txt"
    src.write_text(result.text, encoding="utf-8")
    convert_many([str(src)], tmp_path / "o.docx", "formal", title=result.structure.title, section_titles=False)
    texts = [p.text for p in Document(str(tmp_path / "o.docx")).paragraphs]
    assert texts.count("충원 현황 정리") == 1
