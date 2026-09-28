"""사내 서식 규격이 실제 문서에 그대로 나오는지."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH

from doc2report.pipeline import convert
from doc2report.profile import load_profile
from doc2report.units import emu_to_pt

FIXTURE = Path(__file__).parent / "fixtures" / "company_format.md"


@pytest.fixture(scope="module")
def docx(tmp_path_factory):
    out = tmp_path_factory.mktemp("company") / "report.docx"
    convert(str(FIXTURE), out, "default")
    return DocxDocument(str(out))


@pytest.fixture(scope="module")
def profile():
    return load_profile("default")


def paragraph_at(docx, index: int):
    return [p for p in docx.paragraphs if p.text.strip()][index]


def test_title_is_centered_bold_and_underlined(docx, profile):
    title = paragraph_at(docx, 0)
    spec = profile.font("title")
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert all(run.bold for run in title.runs)
    assert all(run.underline for run in title.runs)
    assert emu_to_pt(int(title.runs[0].font.size)) == pytest.approx(emu_to_pt(spec.size))


def test_dateline_is_right_aligned_with_date_style(docx, profile):
    line = paragraph_at(docx, 1)
    assert line.text.strip() == "2026. 10. 1"  # 2026-10-01 을 프로파일 형식으로
    assert line.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert emu_to_pt(int(line.runs[0].font.size)) == pytest.approx(
        emu_to_pt(profile.font("date").size))


def test_paragraph_hierarchy_is_number_square_dash(docx, profile):
    texts = [p.text for p in docx.paragraphs if p.text.strip()]
    assert any(t.startswith("1.\t") for t in texts)
    assert any(t.startswith("□\t") for t in texts)
    assert any(t.startswith("-\t") for t in texts)


def test_numbering_continues_across_sections(docx):
    numbers = [p.text.split(".")[0] for p in docx.paragraphs
               if p.text.strip() and p.text[0].isdigit() and "\t" in p.text]
    assert numbers == ["1", "2", "3", "4"]


def test_first_two_levels_are_bold_third_is_not(docx):
    for paragraph in docx.paragraphs:
        text = paragraph.text
        if text.startswith(("1.\t", "2.\t", "□\t")):
            assert all(run.bold for run in paragraph.runs), text
        elif text.startswith("-\t"):
            assert not any(run.bold for run in paragraph.runs), text


def test_indent_matches_profile_levels(docx, profile):
    expected = {
        "1.\t": profile.numbering_level(0),
        "□\t": profile.numbering_level(1),
        "-\t": profile.numbering_level(2),
    }
    for paragraph in docx.paragraphs:
        for prefix, level in expected.items():
            if paragraph.text.startswith(prefix):
                # docx는 twips로 저장하므로 1 twip(635 EMU) 반올림 오차를 허용한다.
                left = int(paragraph.paragraph_format.left_indent or 0)
                assert abs(left - (level.indent + (level.hanging or 0))) <= 635
                hanging = int(paragraph.paragraph_format.first_line_indent or 0)
                assert abs(hanging + (level.hanging or 0)) <= 635


def test_spacing_only_when_level_changes(docx, profile):
    """같은 단계가 이어지면 붙이고, 단계가 바뀌는 자리에서만 띄운다."""
    items = [p for p in docx.paragraphs if p.text.startswith(("1.\t", "□\t", "-\t"))]
    spacings = {int(p.paragraph_format.space_after or 0) for p in items}
    change = profile.numbering_level(0).space_after_level_change
    assert change in spacings          # 단계가 바뀌는 자리가 있고
    assert 0 in spacings               # 붙어 있는 자리도 있다


def test_new_section_is_clearly_separated(docx, profile):
    """- 로 끝난 뒤 2. 가 시작되는 자리는 단계 전환 간격보다 확실히 넓어야 한다."""
    sections = [p for p in docx.paragraphs
                if p.text.startswith(("2.\t", "3.\t", "4.\t"))]
    assert sections
    minimum = profile.numbering_level(0).space_after_level_change
    for paragraph in sections:
        before = int(paragraph.paragraph_format.space_before or 0)
        assert before > minimum, paragraph.text


def test_paragraph_after_table_is_pushed_down(docx, profile):
    """표 바로 다음 문단이 표에 붙지 않는다."""
    following = [p for p in docx.paragraphs if p.text.startswith("4.\t")][0]
    assert int(following.paragraph_format.space_before or 0) >= profile.tables.space_after


def test_table_cells_are_centered_and_single_spaced(docx, profile):
    table = docx.tables[0]
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
                assert paragraph.paragraph_format.line_spacing == pytest.approx(
                    profile.font("table").line_spacing)


def test_rows_have_a_minimum_height(docx, profile):
    """행 최소 높이는 기본값, 지면에 여유가 있으면 넉넉한 값 — 둘 중 하나여야 한다."""
    from docx.oxml.ns import qn

    allowed = {round(profile.tables.row_height / 635),
               round(profile.tables.row_height_relaxed / 635)}
    for row in docx.tables[0].rows:
        height = row._tr.find(qn("w:trPr")).find(qn("w:trHeight"))
        assert height is not None
        assert height.get(qn("w:hRule")) == "atLeast"
        assert min(abs(int(height.get(qn("w:val"))) - a) for a in allowed) <= 1


def test_overrides_change_font_and_size(tmp_path, profile):
    out = tmp_path / "override.docx"
    convert(str(FIXTURE), out,
            profile.with_overrides(font="맑은 고딕", size="12pt", line_spacing="1.5"))
    docx = DocxDocument(str(out))

    body = [p for p in docx.paragraphs if p.text.startswith("□\t")][0]
    assert emu_to_pt(int(body.runs[0].font.size)) == pytest.approx(12)
    assert body.paragraph_format.line_spacing == pytest.approx(1.5)
    assert body.runs[0].font.name == "맑은 고딕"


def test_overrides_do_not_touch_the_profile_file(profile):
    """--ask 로 고른 값은 이번 변환에만 적용된다."""
    changed = profile.with_overrides(size="9pt")
    assert changed.font("body").size != profile.font("body").size
    assert load_profile("default").font("body").size == profile.font("body").size
