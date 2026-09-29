"""연결 페이지 한 번에 불러오기 — 본문 페이지 안의 페이지 포함·발췌 포함·하위 페이지·첨부 보기 매크로
(2026-09-29 사용자: "+로 펼쳐 보게 연결해 둔 페이지를 본문 하나로 한 번에 변환").

실제 Confluence 대신 httpx.MockTransport로 REST 응답을 흉내 낸다.
"""

from __future__ import annotations

import io

import httpx
import pytest
from docx import Document as DocxDocument

from doc2report.ir import Callout, Heading, Paragraph, Table, plain
from doc2report.parsers.confluence_storage import PageRef, parse_confluence_storage
from doc2report.pipeline import load_document

MAIN = """
<h2>1. 추진 배경</h2>
<p>본문 문장</p>
<ac:structured-macro ac:name="expand">
  <ac:parameter ac:name="title">세부 계획 펼치기</ac:parameter>
  <ac:rich-text-body>
    <ac:structured-macro ac:name="include">
      <ac:parameter ac:name=""><ac:link><ri:page ri:content-title="세부 계획"/></ac:link></ac:parameter>
    </ac:structured-macro>
  </ac:rich-text-body>
</ac:structured-macro>
<ac:structured-macro ac:name="excerpt-include">
  <ac:parameter ac:name=""><ac:link><ri:page ri:content-title="요약 페이지" ri:space-key="HR"/></ac:link></ac:parameter>
</ac:structured-macro>
<ac:structured-macro ac:name="view-file">
  <ac:parameter ac:name="name"><ri:attachment ri:filename="첨부보고.docx"/></ac:parameter>
</ac:structured-macro>
<ac:structured-macro ac:name="view-file">
  <ac:parameter ac:name="name"><ri:attachment ri:filename="예산.xlsx"/></ac:parameter>
</ac:structured-macro>
<ac:structured-macro ac:name="children"/>
"""

PAGES = {
    "세부 계획": {"id": "200", "title": "세부 계획", "space": {"key": "TEAM"}, "body": {"storage": {"value":
        '<p>□ 세부 항목</p><ac:structured-macro ac:name="include"><ac:parameter ac:name="">'
        '<ac:link><ri:page ri:content-title="본문"/></ac:link></ac:parameter></ac:structured-macro>'}}},
    "요약 페이지": {"id": "300", "title": "요약 페이지", "space": {"key": "HR"}, "body": {"storage": {"value":
        '<p>발췌 밖 문장</p><ac:structured-macro ac:name="excerpt"><ac:rich-text-body>'
        '<p>발췌 안 문장</p></ac:rich-text-body></ac:structured-macro>'}}},
    "본문": {"id": "100", "title": "본문", "space": {"key": "TEAM"}, "body": {"storage": {"value": MAIN}}},
}
CHILDREN = [{"id": "400", "title": "하위 페이지 A", "space": {"key": "TEAM"},
             "body": {"storage": {"value": "<p>하위 A 내용</p>"}}}]


def _docx_bytes() -> bytes:
    d = DocxDocument()
    d.add_paragraph("첨부 Word 문장")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def confluence(monkeypatch):
    monkeypatch.setenv("CONFLUENCE_URL", "https://wiki.company.com")
    monkeypatch.setenv("CONFLUENCE_API_TOKEN", "pat")
    seen: list[str] = []
    docx = _docx_bytes()

    def handler(request: httpx.Request) -> httpx.Response:
        path, params = request.url.path, request.url.params
        seen.append(f"{path}?{params}")
        if path == "/rest/api/content/100":
            return httpx.Response(200, json=PAGES["본문"])
        if path == "/rest/api/content":
            page = PAGES.get(params.get("title"))
            return httpx.Response(200, json={"results": [page] if page else []})
        if path == "/rest/api/content/100/child/page":
            return httpx.Response(200, json={"results": CHILDREN})
        if path == "/rest/api/content/100/child/attachment":
            return httpx.Response(200, json={"results": [
                {"title": "첨부보고.docx", "_links": {"download": "/download/att.docx"}},
                {"title": "예산.xlsx", "_links": {"download": "/download/b.xlsx"}}]})
        if path.endswith("/child/attachment"):
            return httpx.Response(200, json={"results": []})
        if path == "/download/att.docx":
            return httpx.Response(200, content=docx)
        if path == "/download/b.xlsx":
            return httpx.Response(200, content=b"xlsx")
        return httpx.Response(404, json={})

    real = httpx.Client

    def fake_client(*args, **kwargs):
        kwargs["transport"] = httpx.MockTransport(handler)
        return real(*args, **kwargs)

    monkeypatch.setattr("doc2report.sources.confluence.httpx.Client", fake_client)
    return seen


def _texts(blocks) -> list[str]:
    out = []
    for block in blocks:
        if isinstance(block, (Paragraph, Heading)):
            out.append(plain(block.runs))
        elif isinstance(block, Callout):
            out.extend(_texts(block.blocks))
    return out


def test_linked_pages_are_loaded_into_one_document(confluence):
    doc, notes = load_document("https://wiki.company.com/pages/viewpage.action?pageId=100")
    texts = _texts(doc.blocks)
    assert texts.index("세부 계획 펼치기") < texts.index("□ 세부 항목")  # 펼치기 제목 + 포함 페이지
    assert "발췌 안 문장" in texts and "발췌 밖 문장" not in texts   # 발췌 포함은 excerpt만
    assert "첨부 Word 문장" in texts                                   # 첨부 Word는 내용째
    assert "※ 첨부: 예산.xlsx" in texts                                 # 그 밖의 첨부는 이름만
    assert texts.index("하위 페이지 A") < texts.index("하위 A 내용")      # 하위 페이지는 제목 + 내용
    assert any(isinstance(b, Heading) and plain(b.runs) == "하위 페이지 A" for b in doc.blocks)
    assert texts.count("본문 문장") == 1                                # 세부 계획 → 본문 순환은 막음
    assert any("순환" in n for n in notes)
    assert any("'세부 계획'를 함께 불러옴" in n for n in notes)
    assert any("spaceKey=HR" in s for s in confluence)                  # 다른 스페이스 지정도 따름
    assert not any(isinstance(b, PageRef) for b in doc.blocks)


def test_linked_pages_can_be_turned_off(confluence):
    doc, notes = load_document("https://wiki.company.com/pages/100", linked=False)
    texts = _texts(doc.blocks)
    assert "□ 세부 항목" not in texts and "세부 계획 펼치기" in texts
    assert any("포함 페이지 '세부 계획'을(를) 불러오지 않음" in n for n in notes)
    assert not any("/rest/api/content?" in s for s in confluence)


def test_parser_without_resolver_leaves_no_placeholders():
    parsed = parse_confluence_storage(MAIN)
    assert not any(isinstance(b, PageRef) for b in parsed.document.blocks)
    assert any("불러오지 않음" in n for n in parsed.notes)


def test_include_inside_table_cell_leaves_only_the_name(confluence, monkeypatch):
    table_page = ('<table><tbody><tr><td><ac:structured-macro ac:name="include"><ac:parameter ac:name="">'
                  '<ac:link><ri:page ri:content-title="세부 계획"/></ac:link></ac:parameter>'
                  '</ac:structured-macro></td></tr></tbody></table>')
    monkeypatch.setitem(PAGES, "본문", {**PAGES["본문"], "body": {"storage": {"value": table_page}}})
    doc, _ = load_document("https://wiki.company.com/pages/100")
    table = next(b for b in doc.blocks if isinstance(b, Table))
    assert plain(table.rows[0].cells[0].blocks[0].runs) == "(포함 페이지 '세부 계획')"


def test_pagetree_of_whole_space_is_not_loaded():
    macro = '<ac:structured-macro ac:name="pagetree"><ac:parameter ac:name="root">@home</ac:parameter></ac:structured-macro>'
    parsed = parse_confluence_storage(macro, keep_refs=True)
    assert parsed.document.blocks == []
    assert any("스페이스 전체" in n for n in parsed.notes)


INCLUDE = ('<ac:structured-macro ac:name="include"><ac:parameter ac:name="">'
           '<ac:link><ri:page ri:content-title="세부 계획"/></ac:link></ac:parameter></ac:structured-macro>')


@pytest.mark.parametrize("xhtml", [
    f"<p>앞 문장 {INCLUDE} 뒤 문장</p>",           # 편집기가 매크로를 문단 안에 넣은 경우
    f"<ul><li>항목 {INCLUDE}</li></ul>",           # 목록 항목 안
    f"<ul><li><p>항목</p>{INCLUDE}</li></ul>",
])
def test_include_nested_in_paragraph_or_list_is_not_lost(xhtml):
    blocks = parse_confluence_storage(xhtml, keep_refs=True).document.blocks
    refs = [b for b in blocks if isinstance(b, PageRef)]
    assert [r.title for r in refs] == ["세부 계획"]


def test_anchor_bookmark_and_status_do_not_leak_parameters():
    """책갈피(anchor)의 이름·상태(status)의 색 같은 매개변수가 본문 글자로 새어 나오면 안 된다."""
    xhtml = ('<h2><ac:structured-macro ac:name="anchor"><ac:parameter ac:name="">sec1</ac:parameter>'
             '</ac:structured-macro>1. 추진 배경</h2>'
             '<ac:structured-macro ac:name="anchor"><ac:parameter ac:name="">x</ac:parameter></ac:structured-macro>'
             '<p><ac:structured-macro ac:name="status"><ac:parameter ac:name="colour">Green</ac:parameter>'
             '<ac:parameter ac:name="title">완료</ac:parameter></ac:structured-macro> 상태</p>'
             '<table><tbody><tr><td>진행 <ac:structured-macro ac:name="status">'
             '<ac:parameter ac:name="title">지연</ac:parameter></ac:structured-macro></td></tr></tbody></table>'
             '<p><ac:link><ri:page ri:content-title="세부 계획"/></ac:link></p>')
    parsed = parse_confluence_storage(xhtml)
    texts = _texts(parsed.document.blocks)
    assert texts == ["1. 추진 배경", "완료 상태", "세부 계획"]
    table = next(b for b in parsed.document.blocks if isinstance(b, Table))
    assert plain(table.rows[0].cells[0].blocks[0].runs) == "진행 지연"   # 예전엔 칸이 통째로 비었다
    assert parsed.notes == []   # 책갈피는 "지원하지 않는 매크로"가 아니다
