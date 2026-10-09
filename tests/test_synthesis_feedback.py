"""종합 실측(2026-10-08, 사내 LLM A·B) 사용자 피드백 반영."""

import json

import pytest
from docx import Document as OpenDocx

from doc2report.drafting import Line, Rewrite, SynthSpec, fix_change_arrows, review
from doc2report.profile import load_profile
from doc2report.render import docx_writer
from doc2report.synthesis import SYNTH_RULES, prepare
from doc2report.transform.factcheck import check, load_rules

from test_formal_lines import fake_fonts  # noqa: F401
from test_synthesis import DOC1, DOC2, _doc, _fake


def test_spaced_report_dates_count_as_source_dates():
    """B: 원문 '2026. 7. 1 ~ 8. 29'를 '7.1 ~ 8.29'로 쓴 줄이 '원문에 없는 날짜'로 걸려 원문풍 49자 줄로 바뀌었다."""
    src = "기간 : 2026. 7. 1 ~ 8. 29 (8주)"
    assert check("운영 기간 : 7.1 ~ 8.29 (8주)", src, year=2026) == []
    assert check("운영 기간 : 7.2 ~ 8.29 (8주)", src, year=2026)             # 틀린 날짜는 여전히 잡는다
    assert check("1. 현황 3. 2", "1. 현 황 3개", year=2026)                  # 번호 "1. 현황"을 날짜로 보지 않는다


@pytest.mark.parametrize("before, after", [
    ("리드타임  : 28주 (기존 26주 → 8월 전력 제한 영향)", "리드타임  : 기존 26주 → 28주 (8월 전력 제한 영향)"),
    ("- 3,900만원 (당초 4,200만원 → 협상 결과)", "- 당초 4,200만원 → 3,900만원 (협상 결과)"),
])
def test_arrow_goes_between_two_values(before, after):
    line = Line("-", before, [1])
    assert fix_change_arrows([line]) and line.text == after


@pytest.mark.parametrize("text", ["헤드헌팅 수수료 상향 (기존 25% → 30%)", "입주 : 11월 초 확정 (기존 10월 중순 예상)",
                                  "재고 : 180개 (기존 210개 → 180개)"])
def test_correct_change_notation_is_left_alone(text):
    line = Line("-", text, [1])
    assert fix_change_arrows([line]) == [] and line.text == text


def test_example_in_prompt_uses_value_arrow_value():
    assert "(당초 제시" not in SYNTH_RULES and "기존 4,200만원 → 3,900만원 (협상 결과)" in SYNTH_RULES
    assert "단기 대응" in SYNTH_RULES and "검토 의견" in SYNTH_RULES


def test_synthesis_prompt_example_lines_pass_fact_check():
    """지시문 예시가 우리 사실 검증을 통과해야 LLM이 따라 해도 걸리지 않는다."""
    example = SYNTH_RULES[SYNTH_RULES.index("[종합 예시]"):]
    sentences = {int(n): s for n, s in __import__("re").findall(r"\[(\d+)\] (.+)", example)}
    data = json.loads(example[example.index("{"):example.rindex("}") + 1])
    rules = load_rules()
    for line in data["lines"]:
        text = line.get("text") or " ".join(c for r in line.get("rows", []) for c in r)
        assert check(text, " ".join(sentences[i] for i in line["src"]), rules, 2026) == [], text


KEY = DOC2 + "3. 향후 계획\n □ 중기 목표 : A사 60%, 대체 40%\n"


def test_key_section_topics_are_tracked_and_demanded():
    prep = prepare([_doc(DOC1, "d1"), _doc(KEY, "d2")])
    keyed = [prep.sentences[g[0] - 1] for g in prep.key_groups]
    assert any("중기 목표" in s for s in keyed) and not any("200개" in s for s in keyed if "리드" in s)
    rewrite = Rewrite("t", [Line("□", "리드타임 : 28주", [next(i for i, s in enumerate(prep.sentences, 1) if "28주" in s)])], [])
    _, style = review(rewrite, prep.sentences, "t", 2026, synth=SynthSpec(0, 0, prep.key_groups, prep.number_ids))
    assert any("통째로 빠짐" in s for s in style)


def test_synthesis_keeps_three_candidate_comparison_table():
    """A: 후보 3곳을 같은 항목으로 비교한 표가 ①②③ 줄로 바뀌었다 — 종합에서는 표를 둔다."""
    docs = [_doc(DOC1, "d1"), _doc(DOC2, "d2")]
    prep = prepare(docs)
    ids = list(range(1, len(prep.sentences) + 1))
    reply = json.dumps({"title": "부품 공급 대응", "lines": [
        {"m": "1.", "text": "후보 비교", "src": ids[:1]},
        {"m": "표", "rows": [["구분", "(1안) 갑사", "(2안) 을사", "(3안) 병사"],
                            ["리드타임", "- 26주", "- 26주", "- 28주"],
                            ["재고", "- 2.5주분", "- 2.1주분", "- 2.1주분"]], "src": ids}],
        "dropped": []}, ensure_ascii=False)
    from doc2report.synthesis import synthesize
    result = synthesize(docs, ask=_fake(reply), year=2026)
    assert "\t" in result.text and "①" not in result.text
    assert not any("①②③ 줄로 바꿈" in n for n in result.notes)


def test_list_cells_are_left_aligned_in_formal_tables(tmp_path):
    from doc2report.pipeline import convert
    src = tmp_path / "t.md"
    src.write_text("# 표\n\n| 구분 | (1안) 외부 | (2안) 사내 |\n|---|---|---|\n| 보안 | - 일부 전송 | - 전송 없음 |\n"
                   "| 기간 | - 즉시 | - 약 6개월 |\n", encoding="utf-8")
    out = tmp_path / "o.docx"
    convert(str(src), out, "formal", polish="none")
    table = OpenDocx(str(out)).tables[0]
    aligns = {c.paragraphs[0].alignment for row in table.rows[1:] for c in row.cells[1:]}
    assert aligns == {0}                                    # WD_ALIGN_PARAGRAPH.LEFT


def test_shortening_keeps_aligned_label(tmp_path, fake_fonts):
    """B: '미사용자           : 9명 …'을 LLM이 줄이며 '미사용자: 9명'으로 바꿔 쌍점 맞춤이 깨졌다 — 값만 보낸다."""
    from doc2report.ir import Document, Heading, Paragraph, Run
    from doc2report.layout.table_fit import plan_tables

    prof = load_profile("formal")
    label = "미사용자     : "
    seen: list[str] = []

    def shortener(original, limit):
        seen.append(original)
        return "가나다 가나다"
    for n in range(4, 40):                    # 좁히기 한도로도 두세 글자만 넘쳐 줄임을 시도하는 길이를 찾는다
        text = "□ " + label + " ".join(["가나다"] * n)
        doc = Document(blocks=[Heading(level=2, runs=[Run("배경")]), Paragraph(runs=[Run(text)])])
        out = tmp_path / "o.docx"
        docx_writer.DocxRenderer(prof, plan_tables(doc, prof), shortener=shortener).save(doc, out)
        if seen:
            break
    else:
        pytest.fail("줄임을 시도하는 문장을 못 만듦")
    body = [p.text for p in OpenDocx(str(out)).paragraphs if "가" in p.text]
    assert seen and not seen[0].startswith("미사용자")              # LLM에는 값만
    assert len(body) == 1 and label + "가나다 가나다" in body[0]       # 항목명·맞춤 공백은 그대로


RESULT_DOC = """파일럿 운영 결과
2026. 9. 18
1. 운영 결과
 □ 작업 시간 단축 : 평균 18% (신규 기능 24%, 버그 수정 12%)
   - 테스트 코드 작성 시간은 31% 단축
 □ 사용자 만족도 : 5점 만점 4.1점, 지속 사용 의향 88%
"""
NEXT_DOC = """확산 계획
2026. 10. 5
1. 추진 경과
 □ 파일럿 결과 작업 시간 평균 18% 단축, 만족도 4.1점
2. 확산 단계
 □ 1단계 : 150석 도입
"""


def test_result_topics_must_not_vanish_but_wording_is_left_to_the_llm():
    """B: 파일럿 결과가 뒤 문서의 한 줄 요약만 남고 사라졌다. 규칙은 주제(□ 항목 + 하위)가 통째로 빠지는 것만 막는다 —
    문장 모양(숫자 유무·단계)으로 중요도를 단정하지 않는다(2026-10-09 사용자)."""
    prep = prepare([_doc(RESULT_DOC, "d1"), _doc(NEXT_DOC, "d2")])
    groups = [[prep.sentences[i - 1][:6] for i in g] for g in prep.key_groups]
    assert any(g[0].startswith("작업 시간") and len(g) == 2 for g in groups)       # 하위 "- 테스트 코드 …"가 같은 주제에 묶인다
    assert any(g[0].startswith("사용자 만족") for g in groups)
    idx = {s[:6]: i for i, s in enumerate(prep.sentences, 1)}
    spec = SynthSpec(0, 0, prep.key_groups, prep.number_ids)
    summary = next(i for i, s in enumerate(prep.sentences, 1) if s.startswith("파일럿 결과"))
    thin = Rewrite("t", [Line("□", "파일럿 결과 : 평균 18% 단축", [summary])], [])         # 만족도 주제가 사라짐
    _, style = review(thin, prep.sentences, "t", 2026, synth=spec)
    assert any("주제" in s and "만족" in s for s in style), style
    brief_ok = Rewrite("t", [Line("□", "파일럿 결과", [idx["작업 시간 단축 :"[:6]]]),     # 하위 문장 없이 맨 위 문장만 써도 통과
                             Line("-", "만족도 4.1점", [idx["사용자 만족도"[:6]]])], [])
    _, style = review(brief_ok, prep.sentences, "t", 2026, synth=spec)
    assert not [s for s in style if "주제" in s], style


def test_topics_without_digits_or_short_phrases_are_protected_too():
    """숫자 없는 □ 항목·짧은 구로 된 □ 항목이 핵심인 문서에서도 주제가 빠지지 않게 한다."""
    doc = "검토 보고\n2026. 10. 1\n1. 검토 결과\n □ 보안 우려 해소\n □ 예산 부담\n   - 증액 불가\n"
    prep = prepare([_doc(doc, "d1"), _doc(NEXT_DOC, "d2")])
    texts = [[prep.sentences[i - 1] for i in g] for g in prep.key_groups]
    assert ["보안 우려 해소"] in texts and ["예산 부담", "증액 불가"] in texts
    only_first = Rewrite("t", [Line("□", "보안 우려 해소", [prep.sentences.index("보안 우려 해소") + 1])], [])
    _, style = review(only_first, prep.sentences, "t", 2026, synth=SynthSpec(0, 0, prep.key_groups, prep.number_ids))
    assert any("예산 부담" in s for s in style), style


def test_numbers_in_cited_plan_sentences_must_survive():
    prep = prepare([_doc(DOC1, "d1"), _doc(KEY, "d2")])
    target = next(i for i, s in enumerate(prep.sentences, 1) if "중기 목표" in s)
    assert target in prep.number_ids
    lost = Rewrite("t", [Line("□", "중기 목표 설정", [target])], [])
    _, style = review(lost, prep.sentences, "t", 2026, synth=SynthSpec(0, 0, prep.key_groups, prep.number_ids))
    assert any("수치" in s and "60" in s for s in style), style


def test_synthesis_without_title_gets_a_fallback_title():
    """B 1회: 제목 없이 출력됐다 — review가 빈 제목을 재작성 사유로 올리고, 끝내 비면 본문으로 다시 묻거나 최신 문서 제목을 쓴다."""
    docs = [_doc(DOC1, "d1"), _doc(DOC2, "d2")]
    prep = prepare(docs)
    ids = {s: i for i, s in enumerate(prep.sentences, 1)}
    find = lambda part: next(i for s, i in ids.items() if part in s)   # noqa: E731
    body = {"title": "", "lines": [
        {"m": "1.", "text": "현 황", "src": [find("28주")]},
        {"m": "□", "text": "리드타임 : 기존 26주 → 28주", "src": [find("26주"), find("28주")]}], "dropped": []}
    seen = []

    def ask(system, user):
        seen.append(system[:12])
        if "제목을 한 줄로" in system:
            return "부품 수급 현황 및 대응\n"
        return json.dumps(body, ensure_ascii=False)
    from doc2report.synthesis import synthesize
    result = synthesize(docs, ask=ask, year=2026)
    assert result.title == "부품 수급 현황 및 대응"
    assert any(n.startswith("제목:") for n in result.notes)
    assert sum(1 for s in seen if not s.startswith("보고서 본문")) >= 2      # 빈 제목 때문에 한 번 다시 썼다

    def broken_title(system, user):
        if "제목을 한 줄로" in system:
            raise RuntimeError("down")
        return json.dumps(body, ensure_ascii=False)
    result = synthesize(docs, ask=broken_title, year=2026)
    assert result.title == "부품 대응 계획"                                     # 가장 최근(10. 2) 문서 제목
