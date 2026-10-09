"""Confluence 변환 5건 수정(2026-10-09 사용자): 공백 들여쓰기·빈 줄 보존·표 한 쪽 유지·쪽 채우기."""

from docx import Document as OpenDocx

from doc2report import pipeline
from doc2report.ir import BlankLine, Cell, Document, ListItem, PageBreak, Paragraph, Row, Run, Table
from doc2report.layout.flow import plan_flow
from doc2report.layout.table_fit import plan_tables
from doc2report.parsers.confluence_storage import parse_confluence_storage
from doc2report.profile import load_profile
from doc2report.render.docx_writer import DocxRenderer
from doc2report.transform.structure import settle_blank_lines


def _kinds(doc):
    return [type(b).__name__ for b in doc.blocks]


def test_parser_turns_empty_paragraphs_into_blank_lines():
    doc = parse_confluence_storage("<p>가</p><p><br/></p><p></p><p>나<br/><br/>다</p><p>끝<br/></p>").document
    assert _kinds(doc) == ["Paragraph", "BlankLine", "BlankLine", "Paragraph", "BlankLine", "Paragraph", "Paragraph"]


def test_blank_lines_inside_table_cells_are_not_kept():
    doc = parse_confluence_storage("<table><tbody><tr><td><p>가</p><p><br/></p><p>나</p></td></tr></tbody></table>").document
    cell = doc.blocks[0].rows[0].cells[0]
    assert [type(b).__name__ for b in cell.blocks] == ["Paragraph", "Paragraph"]


def test_settle_collapses_and_trims_blank_lines():
    table = Table(rows=[Row(cells=[Cell(blocks=[Paragraph(runs=[Run("x")])])])])
    doc = Document(blocks=[BlankLine(), Paragraph(runs=[Run("가")]), BlankLine(), BlankLine(), Paragraph(runs=[Run("나")]),
                           table, BlankLine(), Paragraph(runs=[Run("다")]), BlankLine(), PageBreak(), BlankLine()])
    kept = settle_blank_lines(doc, keep=True)
    assert _kinds(kept) == ["Paragraph", "BlankLine", "Paragraph", "Table", "Paragraph", "PageBreak"]
    assert "BlankLine" not in _kinds(settle_blank_lines(doc, keep=False))


def test_blank_line_between_bracket_caption_and_table_is_dropped():
    table = Table(rows=[Row(cells=[Cell(blocks=[Paragraph(runs=[Run("x")])])])])
    doc = Document(blocks=[Paragraph(runs=[Run("【현황】")]), BlankLine(), table])
    assert _kinds(settle_blank_lines(doc, keep=True)) == ["Paragraph", "Table"]


def _convert(tmp_path, xhtml, profile):
    from doc2report.sources import LoadedSource

    src = LoadedSource(text=xhtml, name="https://wiki/pages/9", format="confluence_storage", title="보고")
    original = pipeline.load_source
    pipeline.load_source = lambda source, **kw: src
    try:
        out = tmp_path / "o.docx"
        result = pipeline.convert("https://wiki/pages/9", out, profile=profile)
    finally:
        pipeline.load_source = original
    return result, [p.text for p in OpenDocx(str(out)).paragraphs]


def test_confluence_keeps_blank_line_and_formal_does_not(tmp_path):
    xhtml = "<p>가 문단</p><p><br/></p><p><br/></p><p>나 문단</p>"
    _, texts = _convert(tmp_path, xhtml, load_profile("confluence"))
    assert texts.count("") == 1                      # 연속 빈 줄은 한 줄로
    _, formal = _convert(tmp_path, xhtml, load_profile("formal"))
    assert "" not in formal


def test_levels_are_two_spaces_apart_in_confluence(tmp_path):
    xhtml = "<ul><li>□ 가<ul><li>- 나<ul><li>ㆍ 다</li></ul></li></ul></li></ul><p>※ 참고</p>"
    _, texts = _convert(tmp_path, xhtml, load_profile("confluence"))
    assert [t for t in texts if t][-4:] == ["□ 가", "  - 나", "    ㆍ 다", "      ※ 참고"]


def test_wrapped_lines_hang_under_text_when_width_is_known(tmp_path, monkeypatch):
    class Fake:
        widths_known = True
        widths_estimated = False
        font_available = False

        def width(self, text):
            return len(text) * 100000

        def char_width(self, ch):
            return 100000

    monkeypatch.setattr(DocxRenderer, "_measurer", lambda self, spec, bold: Fake())
    doc = Document(blocks=[ListItem(depth=1, runs=[Run("본문")], marker="-")])
    out = tmp_path / "h.docx"
    DocxRenderer(load_profile("confluence")).save(doc, out)
    para = [p for p in OpenDocx(str(out)).paragraphs if p.text][0]
    assert para.text == "  - 본문"
    assert abs(int(para.paragraph_format.left_indent) - 4 * 100000) < 700      # "  - " 폭 (twips 반올림)
    assert abs(int(para.paragraph_format.first_line_indent) + 4 * 100000) < 700


def test_confluence_profile_keeps_tables_on_page_and_fills_pages():
    prof = load_profile("confluence")
    assert prof.tables.keep_on_page and prof.tables.page_fill_target == 0.8
    assert prof.text.keep_blank_lines and prof.text.hang_after_marker
    assert not load_profile("formal").text.keep_blank_lines


def test_underfull_page_with_table_stretches_rows():
    prof = load_profile("confluence")
    cell = lambda t, h=False: Cell(blocks=[Paragraph(runs=[Run(t)])], is_header=h)
    table = Table(rows=[Row(cells=[cell("구분", True), cell("내용", True)]),
                        Row(cells=[cell("가"), cell("나")]), Row(cells=[cell("다"), cell("라")])])
    doc = Document(blocks=[Paragraph(runs=[Run("도입")]), table])
    layouts = plan_tables(doc, prof)
    flow = plan_flow(doc, prof, layouts)
    assert flow.row_extra.get(id(table), 0) > 0
    assert any("행 높이" in n for n in flow.notes)
    cap = prof.tables.row_stretch_max
    assert flow.row_extra[id(table)] <= cap
    # 렌더러가 실제 행 최소 높이에 반영한다
    out_doc = DocxRenderer(prof, layouts, flow)
    from pathlib import Path
    import tempfile
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "t.docx"
        out_doc.save(doc, path)
        rows = OpenDocx(str(path)).tables[0].rows
        heights = [r.height for r in rows]
    assert all(h and int(h) >= int(prof.tables.min_row_height(False) or 0) for h in heights)
    assert max(int(h) for h in heights) > int(prof.tables.min_row_height(False) or 0)


def test_profile_round_trips_new_fields(tmp_path):
    from doc2report.profile import dump_profile

    prof = load_profile("confluence")
    path = tmp_path / "p.yaml"
    path.write_text(dump_profile(prof), encoding="utf-8")
    again = load_profile(path)
    assert again.tables.page_fill_target == prof.tables.page_fill_target
    assert again.tables.row_stretch_max == prof.tables.row_stretch_max
    assert again.text.note_extra_spaces == 2 and again.text.keep_blank_lines
