"""제목 접기(fold_headings_into_levels) — 특히 Confluence 문서에 흔한, 제목에
이미 번호가 박혀 있는 경우("## 1. 추진 배경")를 다룬다.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument

from doc2report.ir import Document, Heading, ListItem, Paragraph, Run, Table
from doc2report.pipeline import convert
from doc2report.profile import load_profile
from doc2report.transform.structure import (
    attach_table_captions,
    fold_headings_into_levels,
    merge_short_list_items,
)

FIXTURE = Path(__file__).parent / "fixtures" / "sample_report.md"


@pytest.fixture
def profile():
    return load_profile("default")


def test_existing_heading_number_is_not_duplicated():
    """"## 1. 추진 배경"처럼 제목에 이미 번호가 있으면 프로파일이 매기는 말머리와
    겹치지 않도록 원래 번호를 뗀다 ("1.\\t1. 추진 배경"이 아니라 "1.\\t추진 배경")."""
    doc = Document(blocks=[Heading(level=2, runs=[Run("1. 추진 배경")])])
    folded, changes = fold_headings_into_levels(doc)
    assert folded.blocks[0].runs[0].text == "추진 배경"
    assert len(changes) == 1
    assert changes[0].before == "1. 추진 배경"
    assert changes[0].after == "추진 배경"


def test_heading_without_number_is_untouched():
    doc = Document(blocks=[Heading(level=2, runs=[Run("추진 배경")])])
    folded, changes = fold_headings_into_levels(doc)
    assert folded.blocks[0].runs[0].text == "추진 배경"
    assert changes == []


def test_sample_report_headings_are_not_double_numbered(tmp_path, profile):
    out = tmp_path / "sample.docx"
    convert(str(FIXTURE), out, profile)
    docx = DocxDocument(str(out))
    numbered = [p.text for p in docx.paragraphs if p.text.startswith(("1.\t", "2.\t", "3.\t", "4.\t"))]
    assert numbered, "번호 매겨진 단계가 있어야 한다"
    for text in numbered:
        marker, _, rest = text.partition("\t")
        assert not rest.startswith(marker), f"제목 번호가 겹침: {text!r}"


@pytest.mark.parametrize("text", ["1.1. 추진 배경", "1.추진 배경", "(1) 추진 배경", "2) 추진 배경"])
def test_other_manual_number_styles_are_stripped(text):
    doc = Document(blocks=[Heading(level=2, runs=[Run(text)])])
    folded, _ = fold_headings_into_levels(doc)
    assert folded.blocks[0].runs[0].text == "추진 배경"


@pytest.mark.parametrize("text", ["1.5배 향상 방안", "2026년 계획", "3D 설계"])
def test_headings_that_merely_start_with_a_number_are_kept(text):
    doc = Document(blocks=[Heading(level=2, runs=[Run(text)])])
    folded, changes = fold_headings_into_levels(doc)
    assert folded.blocks[0].runs[0].text == text
    assert changes == []


def test_existing_bullet_marker_in_heading_is_not_duplicated():
    """Confluence 원문 제목에 이미 "□ "가 박혀 있으면(사내 관행) 접을 때 프로파일이
    매기는 말머리와 겹쳐 "ㅁㅁ 채용진행현황"처럼 보이면 안 된다 (2026-09-29 실제 변환)."""
    doc = Document(blocks=[Heading(level=2, runs=[Run("□ 채용진행현황")])])
    folded, changes = fold_headings_into_levels(doc, ["□", "-", "·", "."])
    assert folded.blocks[0].runs[0].text == "채용진행현황"
    assert len(changes) == 1


def test_existing_bullet_marker_in_list_item_is_not_duplicated():
    doc = Document(blocks=[ListItem(depth=0, runs=[Run("- 입사확정")])])
    folded, changes = fold_headings_into_levels(doc, ["□", "-", "·", "."])
    assert folded.blocks[0].runs[0].text == "입사확정"
    assert len(changes) == 1


def test_leading_negative_number_is_not_mistaken_for_a_dash_marker():
    """"-5%" 처럼 마커 뒤에 공백이 없으면 음수로 보고 그대로 둔다."""
    doc = Document(blocks=[ListItem(depth=0, runs=[Run("-5%p 개선")])])
    folded, changes = fold_headings_into_levels(doc, ["□", "-", "·", "."])
    assert folded.blocks[0].runs[0].text == "-5%p 개선"
    assert changes == []


def test_without_strip_markers_configured_nothing_is_stripped():
    doc = Document(blocks=[Heading(level=2, runs=[Run("□ 채용진행현황")])])
    folded, changes = fold_headings_into_levels(doc)
    assert folded.blocks[0].runs[0].text == "□ 채용진행현황"
    assert changes == []


def test_bracket_caption_above_table_is_moved_to_table_caption():
    """표 위 "[사업현황]" 같은 꺾쇠 문단은 Table.caption으로 옮겨 제목 접기에서
    "- [사업현황]"처럼 말머리가 붙지 않게 한다 (2026-09-29 사용자 요청)."""
    doc = Document(blocks=[Paragraph(runs=[Run("[사업현황]")]), Table(rows=[])])
    result, changes = attach_table_captions(doc)
    assert len(result.blocks) == 1
    assert isinstance(result.blocks[0], Table)
    assert result.blocks[0].caption == "사업현황"
    assert len(changes) == 1


def test_paragraph_without_brackets_is_not_treated_as_caption():
    doc = Document(blocks=[Paragraph(runs=[Run("평범한 문단")]), Table(rows=[])])
    result, changes = attach_table_captions(doc)
    assert len(result.blocks) == 2
    assert result.blocks[1].caption is None
    assert changes == []


def test_quote_under_table_becomes_small_star_note(tmp_path, profile):
    """표 아래 설명(인용문)은 '* '로 시작하는 표 주석이 되고, 글자는 table_note 크기."""
    out = tmp_path / "sample.docx"
    convert(str(FIXTURE), out, profile)
    docx = DocxDocument(str(out))
    note = next(p for p in docx.paragraphs if "측정 기준은" in p.text)
    assert note.text.startswith(f"{profile.tables.note_marker} ")
    assert all(int(run.font.size) == profile.font("table_note").size for run in note.runs)
    # 주석은 표에 붙어 있고, '표 뒤 간격'은 그 다음 항목이 가져간다.
    assert int(note.paragraph_format.space_before or 0) < profile.tables.space_after
    following = next(p for p in docx.paragraphs if p.text.startswith("4.\t"))
    assert int(following.paragraph_format.space_before or 0) >= profile.tables.space_after


def test_heading_ending_in_a_copula_becomes_a_bare_noun(tmp_path, profile):
    """제목은 명사로 끝나야 한다("추진 배경임" → "추진 배경", 2026-09-28 사용자 요청)."""
    source = tmp_path / "heading.md"
    source.write_text("# 보고서\n\n## 추진 배경임\n\n본문.\n", encoding="utf-8")
    out = tmp_path / "heading.docx"
    convert(str(source), out, profile)
    docx = DocxDocument(str(out))
    assert any(p.text == "1.\t추진 배경" for p in docx.paragraphs)
    assert not any("배경임" in p.text for p in docx.paragraphs)


def test_short_adjacent_items_are_merged_with_and(tmp_path, profile):
    """사용자 요청(2026-09-28): 짧은 항목끼리는 "및"으로 한 줄에 합친다."""
    out = tmp_path / "sample.docx"
    convert(str(FIXTURE), out, profile)
    docx = DocxDocument(str(out))
    texts = [p.text for p in docx.paragraphs]
    assert any(t == "□\t10월 중 2차 성능 시험 실시 및 미흡 사항은 4분기 과제로 이관"
              for t in texts)
    # 셋째 항목은 짝이 없어 그대로 남는다.
    assert any(t == "□\t결과는 월간 운영보고에 반영 예정" for t in texts)


def test_merge_short_list_items_from_ir():
    from doc2report.ir import ListItem

    doc = Document(blocks=[
        ListItem(depth=1, runs=[Run("짧은 항목 1")]),
        ListItem(depth=1, runs=[Run("짧은 항목 2")]),
    ])
    merged, changes = merge_short_list_items(doc, max_chars=60)
    assert len(merged.blocks) == 1
    assert merged.blocks[0].runs[0].text == "짧은 항목 1 및 짧은 항목 2"
    assert changes[0].rule == "항목 병합"


def test_merge_never_chains_three_items():
    """A+B+C가 모두 짧아도 A+B, C(홀로)까지만 — "및"이 두 번 겹치지 않는다."""
    from doc2report.ir import ListItem

    doc = Document(blocks=[
        ListItem(depth=1, runs=[Run("가")]),
        ListItem(depth=1, runs=[Run("나")]),
        ListItem(depth=1, runs=[Run("다")]),
    ])
    merged, _ = merge_short_list_items(doc, max_chars=60)
    texts = [b.runs[0].text for b in merged.blocks]
    assert texts == ["가 및 나", "다"]
    assert all(t.count("및") <= 1 for t in texts)


def test_merge_respects_depth_and_length_budget():
    from doc2report.ir import ListItem

    doc = Document(blocks=[
        ListItem(depth=0, runs=[Run("짧음")]),
        ListItem(depth=1, runs=[Run("다른 깊이라 안 합쳐짐")]),  # depth 다름
        ListItem(depth=1, runs=[Run("가" * 60)]),  # 길이 초과라 안 합쳐짐
        ListItem(depth=1, runs=[Run("나")]),
    ])
    merged, changes = merge_short_list_items(doc, max_chars=60)
    assert len(merged.blocks) == 4  # 아무것도 안 합쳐짐
    assert changes == []


def test_merge_is_a_noop_when_threshold_is_zero():
    from doc2report.ir import ListItem

    doc = Document(blocks=[ListItem(depth=1, runs=[Run("가")]),
                           ListItem(depth=1, runs=[Run("나")])])
    merged, changes = merge_short_list_items(doc, max_chars=0)
    assert len(merged.blocks) == 2
    assert changes == []


def test_star_paragraph_under_table_is_a_note():
    from doc2report.ir import Cell, Paragraph, Row, Table
    from doc2report.transform.structure import attach_table_notes

    table = Table(rows=[Row(cells=[Cell(blocks=[Paragraph(runs=[Run("a")])])])])
    doc = Document(blocks=[table, Paragraph(runs=[Run("※ 단위: 천원")]),
                           Paragraph(runs=[Run("본문")])])
    attached, changes = attach_table_notes(doc, ["*", "※"], "*")
    assert [len(b.notes) if isinstance(b, Table) else None for b in attached.blocks] == [1, None]
    assert table.notes[0][0].text == "단위: 천원"
    assert changes[0].after == "* 단위: 천원"
