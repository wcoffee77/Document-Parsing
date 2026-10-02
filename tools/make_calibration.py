"""줄 폭 보정용 docx — Word에서 열어 '한 줄에 몇 글자가 들어가는지' 세어 알려 주는 용도.

    uv run python tools/make_calibration.py [프로파일]    # 기본 formal → out/calibration.docx

각 시험은 라벨 한 줄 + 긴 문장 한 문단이다. 문장이 첫 줄에서 몇 글자까지 들어갔는지(공백 포함)
세어 라벨 옆 번호와 함께 알려 주면 된다. 서식(글꼴·크기·여백)은 프로파일 값을 쓰고,
글자 간격은 도구가 실제로 쓰는 방식(oxml.set_char_spacing)으로 넣는다.
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Emu, Pt  # noqa: E402

from doc2report.profile import load_profile  # noqa: E402
from doc2report.render import oxml  # noqa: E402
from doc2report.render.base_template import open_base_template  # noqa: E402

HANGUL = "가나다라마바사아자차카타파하" * 6
TESTS = [
    # (라벨, 문장, 굵게, 간격 pt)
    *[(f"한글 보통 간격 {c}pt", HANGUL, False, c) for c in (0, 0.2, 0.4, 0.7, 1.0)],
    *[(f"한글 굵게 간격 {c}pt", HANGUL, True, c) for c in (0, 0.4, 1.0)],
    *[(f"한글+공백(5자마다) 간격 {c}pt", " ".join(HANGUL[i:i + 5] for i in range(0, 70, 5)), False, c)
      for c in (0, 1.0)],
    *[(f"한글+영문숫자(5자마다 'ab12') 간격 {c}pt", "".join(HANGUL[i:i + 5] + "ab12" for i in range(0, 70, 5)),
       False, c) for c in (0, 1.0)],
]


def main() -> None:
    name = sys.argv[1] if len(sys.argv) > 1 else "formal"
    profile = load_profile(name)
    doc = Document(profile.template or open_base_template())  # python-docx 기본 .docx 파일은 사내 보안이 막는다(CLAUDE.md)
    if profile.text.balance_sbcs_dbcs:  # 변환기와 같은 호환 옵션(없으면 한글 간격 −1pt에서 37자, 있으면 40자)
        oxml.set_compat_flag(doc.settings.element, "balanceSingleByteDoubleByteWidth")
    section = doc.sections[0]
    page = profile.page
    section.page_width, section.page_height = Emu(page.width), Emu(page.height)
    section.left_margin, section.right_margin = Emu(page.margin.left), Emu(page.margin.right)
    section.top_margin, section.bottom_margin = Emu(page.margin.top), Emu(page.margin.bottom)
    body = profile.font("body")
    for number, (label, text, bold, spacing) in enumerate(TESTS, 1):
        head = doc.add_paragraph()
        head_run = head.add_run(f"[{number}] {label}")
        head_run.font.size = Pt(10)
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        run = paragraph.add_run(text)
        rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("ascii", "hAnsi", "eastAsia"):
            rfonts.set(qn(f"w:{attr}"), body.east_asia)
        run.font.size = Emu(body.size)
        run.bold = bold
        if spacing:
            oxml.set_char_spacing(run, -round(spacing * 20))
    out = Path(__file__).resolve().parents[1] / "out" / "calibration.docx"
    out.parent.mkdir(exist_ok=True)
    doc.save(out)
    print(f"저장: {out} (글꼴 {body.east_asia}, {body.size / 12700:g}pt)")


if __name__ == "__main__":
    main()
