"""정식보고서 .docx에서 '실제로 보이는 서식'과 형식 요소를 그대로 기록한다.

`parsers/docx_reader.py`는 변환용 입력 파서라 빈 문단·서식 값·머리말/꼬리말·텍스트 상자를
버린다(IR 계약). 문서의 **특징을 파악**하려면 그것들이 바로 재료라서, 분석용으로 따로 읽는다.
변환 코드는 건드리지 않는다.

- 값은 판단 없이 '실효값'으로 푼다: 직접 서식 > 번호 수준 > 스타일 체인 > docDefaults.
- OOXML을 아는 코드라 설계 원칙 3에 따라 이 모듈에 가둔다 (분석 쪽 stats.py는 XML을 모른다).
- 여기서 나온 값은 코드에 굳히지 않는다 — 분포로 요약해 `profiles/*.yaml` **후보**로만 낸다.
- 원문 글자는 메모리에만 두고 `to_dict()`는 기본으로 뺀다(기밀 문서가 리포트로 새지 않게).
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path

from docx import Document as open_docx
from docx.oxml.ns import qn

from ..ir import invisible_codes, is_blank
from ..units import EMU_PER_MM, PAGE_SIZES, parse_length

_TWIP_PER_MM = 1440 / 25.4
_PUA = re.compile("[\ue000-\uf8ff]")
# 글쓴이가 문자로 쳐 둔 말머리. 기호는 붙여 써도("ㆍ입사예정") 말머리지만, 글자로도 흔한 것
# (ㅇ - * 번호 가나다)은 뒤에 공백·탭이 있을 때만 말머리로 본다("-5%p", "1.5배", "ㅇㅇ팀" 제외).
_TYPED_MARKER = re.compile(
    r"^(?P<lead>[ \u3000\t]*)"
    r"(?:(?P<sym>[□■○●◦▪▫◆◇▶▷▲△※•‧·ㆍ\u119e➢➤►✓]|\([0-9]{1,2}\)|\([가-힣a-zA-Z]\)|[①-⑳❶-❿⓵-⓾])"
    r"|(?P<wordy>ㅇ|[-–—*]|[0-9]{1,2}[.)]|[가-힣a-zA-Z][.)]|[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ][.)]?))"
    r"(?P<sep>[ \t\u3000]?)"
)
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
_WPS = "{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_VML = "{urn:schemas-microsoft-com:vml}"
_VML_SHAPES = {"shape", "rect", "roundrect", "oval", "line"}


# ── 결과 자료형 ─────────────────────────────────────────────────────────


@dataclass
class Fmt:
    """글자 서식(문단 대표값). None = 어디에도 안 적힘(= Word 기본)."""
    east_asia: str | None = None
    ascii: str | None = None
    size_pt: float | None = None
    bold: bool | None = None
    italic: bool | None = None
    char_scale_pct: int | None = None
    color: str | None = None         # RRGGBB, 테마색은 "theme:accent1", 자동이면 None
    underline: bool | None = None


@dataclass
class ParaProbe:
    where: str                       # body | table | header | footer | textbox
    style: str | None
    text: str
    text_len: int
    blank: bool
    invisible_codes: list[str]
    align: str | None
    left_mm: float
    first_line_mm: float             # 음수 = 내어쓰기
    space_before_pt: float
    space_after_pt: float
    line_rule: str                   # auto | exact | atLeast
    line_value: float                # auto면 배수, 아니면 pt
    fmt: Fmt
    heading_level: int | None
    marker_kind: str | None          # typed | auto | None
    marker: str | None               # typed면 글쓴이가 친 말머리, auto면 번호 정의 원문(lvlText)
    marker_sep: str | None           # tab | space | none (typed 말머리 뒤)
    marker_level: int | None         # auto면 ilvl
    leading_spaces: int              # 들여쓰기를 공백으로 친 개수
    pua_bullet: bool                 # Symbol/Wingdings 글머리(문자로는 못 읽음)
    has_object: bool                 # 그림·도형·텍스트 상자가 붙은 문단(글이 없어도 '빈 줄'이 아니다)
    page_break_before: bool
    has_page_break: bool
    text_runs: int
    box: int | None = None           # where == "textbox"이면 DocProbe.textboxes의 번호
    bold_pattern: str = "none"       # 말머리 뒤 글자 기준: all | none | prefix | suffix | mixed
    underline: bool = False          # 말머리 뒤 글자의 과반이 밑줄
    leading_wide: int = 0            # 앞 공백 중 전각 공백(U+3000) 개수 — 반각 두 칸과 폭이 다르다
    cell_row: int | None = None      # where == "table"이면 칸 위치
    cell_col: int | None = None


@dataclass
class TextBoxProbe:
    """텍스트 상자 하나(주석 상자). 글 자체는 여기 없고 글 서식은 where == "textbox" 문단에 있다."""
    kind: str                        # drawingml | vml | unknown
    floating: bool                   # False = 글자처럼 줄 안에 놓임
    h_rel: str | None = None
    h_offset_mm: float | None = None
    h_align: str | None = None
    v_rel: str | None = None
    v_offset_mm: float | None = None
    v_align: str | None = None
    width_mm: float | None = None
    height_mm: float | None = None
    wrap: str | None = None
    behind: bool | None = None       # DrawingML behindDoc (글 뒤로)
    x_mm: float | None = None        # 쪽 왼쪽 끝에서 상자 왼쪽 끝까지(계산 가능할 때)
    border: str = "unspecified"      # none | line | unspecified
    border_color: str | None = None
    fill: str = "unspecified"        # none | RRGGBB | unspecified
    paragraphs: int = 0
    text_len: int = 0
    anchor_where: str | None = None  # 상자가 붙은 문단의 위치(body | table)
    anchor_marker: str | None = None # 붙은 문단의 말머리
    anchor_text_len: int = 0
    placement: str = ""              # "가로 구역 / 세로 기준" 요약


@dataclass
class PageProbe:
    width_mm: float
    height_mm: float
    top_mm: float
    bottom_mm: float
    left_mm: float
    right_mm: float
    header_mm: float
    footer_mm: float
    orientation: str
    paper: str | None
    title_page: bool                 # 첫 쪽 머리말·꼬리말이 다름


@dataclass
class HeaderFooterProbe:
    kind: str                        # header | footer
    section: int
    text_len: int
    text: str
    has_page_field: bool
    align: str | None
    linked: bool


@dataclass
class TableProbe:
    rows: int
    cols: int
    col_widths_mm: list[float]
    jc: str | None
    header_rows: int
    header_fill: str | None
    merged_cells: int
    cell_margin_mm: dict[str, float]
    nested: bool
    style_id: str | None = None
    style_first_row_fill: str | None = None   # 표 스타일의 '첫 행' 조건부 서식(직접 음영이 없을 때 머리 음영이 여기 있다)
    style_first_row_bold: bool | None = None
    look_first_row: bool | None = None
    total_width_mm: float = 0.0
    indent_mm: float = 0.0
    borders: dict[str, str] = field(default_factory=dict)   # 변 → "single/0.5pt"
    fill_first_row_share: float | None = None  # 칸 직접 음영 비율: 첫 행 / 첫 열 / 나머지
    fill_first_col_share: float | None = None
    fill_other_share: float | None = None
    fill_colors: list[str] = field(default_factory=list)
    cell_border_share: float | None = None    # 칸 자체 테두리가 있는 칸의 비율
    cell_border_kinds: list[str] = field(default_factory=list)
    before_kind: str = ""            # 표 바로 위 줄의 종류
    after_kind: str = ""             # 표 바로 아래 줄의 종류


@dataclass
class DocProbe:
    doc_id: str
    name: str
    sections: list[PageProbe] = field(default_factory=list)
    paragraphs: list[ParaProbe] = field(default_factory=list)
    tables: list[TableProbe] = field(default_factory=list)
    headers_footers: list[HeaderFooterProbe] = field(default_factory=list)
    textbox_count: int = 0
    textboxes: list[TextBoxProbe] = field(default_factory=list)
    toc: bool = False
    page_breaks: int = 0
    notes: list[str] = field(default_factory=list)

    def to_dict(self, *, keep_text: bool = False, keep_name: bool = False) -> dict:
        data = asdict(self)
        if not keep_name:
            data["name"] = ""
        if not keep_text:
            for para in data["paragraphs"]:
                para["text"] = ""
            for hf in data["headers_footers"]:
                hf["text"] = ""
        return data


# ── 진입점 ──────────────────────────────────────────────────────────────


def probe_docx(path: str | Path, *, doc_id: str | None = None) -> DocProbe:
    path = Path(path)
    docx = open_docx(str(path))
    probe = DocProbe(doc_id=doc_id or path.stem, name=path.name)
    _Prober(docx, probe).run()
    return probe


class _Prober:
    def __init__(self, docx, probe: DocProbe):
        self.docx = docx
        self.probe = probe
        self._para_of: dict = {}   # w:p 요소 → ParaProbe (상자가 붙은 문단을 찾는 용도)
        self._prev_body: ParaProbe | None = None   # 표 바로 위 본문 줄
        self._last_table: TableProbe | None = None  # 아래 줄을 아직 못 채운 표
        self.styles = {
            s.get(qn("w:styleId")): s
            for s in docx.styles.element.findall(qn("w:style"))
        }
        self.defaults = docx.styles.element.find(qn("w:docDefaults"))
        self.default_para_style = next(
            (s for s in self.styles.values()
             if s.get(qn("w:type")) == "paragraph" and s.get(qn("w:default")) in ("1", "true")),
            None,
        )
        try:
            self.numbering = docx.part.numbering_part.element
        except (KeyError, NotImplementedError, AttributeError):
            self.numbering = None

    # ── 전체 ────────────────────────────────────────────────────────────

    def run(self) -> None:
        body = self.docx.element.body
        self._body(body, where="body", nested=False)
        self._sections()
        self._textboxes(body)
        self.probe.toc = self._has_toc(body)

    def _body(self, container, *, where: str, nested: bool) -> None:
        for child in container.iterchildren():
            if child.tag == qn("w:p"):
                self._paragraph(child, where)
            elif child.tag == qn("w:tbl"):
                self._table(child, nested=nested)
            elif child.tag == qn("w:sdt"):
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    self._body(content, where=where, nested=nested)

    # ── 문단 ────────────────────────────────────────────────────────────

    def _paragraph(self, p, where: str) -> ParaProbe:
        text = _text(p)
        style = self._para_style(p)
        ppr = self._effective_ppr(p, style)
        numbering = self._numbering_of(p, style)
        marker_kind = marker = sep = None
        level = None
        pua = False
        leading = 0
        if numbering is not None:
            level, marker, pua = numbering
            marker_kind = "auto"
        else:
            found = _typed_marker(text)
            if found is not None:
                marker_kind = "typed"
                marker, sep = found
                leading = len(text) - len(text.lstrip(" \u3000"))
        spacing = ppr.get("spacing", {})
        line_rule = spacing.get("lineRule", "auto")
        line = _num(spacing.get("line"), 240.0)
        line_value = line / 240 if line_rule == "auto" else line / 20
        ind = ppr.get("ind", {})
        left = _num(ind.get("left", ind.get("start")), 0.0)
        first = _num(ind.get("firstLine"), 0.0) - _num(ind.get("hanging"), 0.0)
        heading = self._heading_level(p, style)
        has_object = any(True for _ in p.iter(qn("w:drawing"), qn("w:pict"), qn("w:object")))
        para = ParaProbe(
            where=where,
            style=_style_name(style),
            text=text,
            text_len=len(text),
            blank=is_blank(text) and not has_object,
            invisible_codes=invisible_codes(text),
            align=ppr.get("jc_value"),
            left_mm=round(left / _TWIP_PER_MM, 2),
            first_line_mm=round(first / _TWIP_PER_MM, 2),
            space_before_pt=round(_num(spacing.get("before"), 0.0) / 20, 2),
            space_after_pt=round(_num(spacing.get("after"), 0.0) / 20, 2),
            line_rule=line_rule,
            line_value=round(line_value, 3),
            fmt=self._dominant_fmt(p, style),
            heading_level=heading,
            marker_kind=marker_kind,
            marker=marker,
            marker_sep=sep,
            marker_level=level,
            leading_spaces=leading,
            pua_bullet=pua,
            has_object=has_object,
            page_break_before="pageBreakBefore" in ppr,
            has_page_break=_has_page_break(p),
            text_runs=len(p.findall(qn("w:r"))),
        )
        offset = 0
        if marker_kind == "typed":
            offset = leading + len(marker or "") + (1 if sep in ("tab", "space") else 0)
        para.bold_pattern, para.underline = self._bold_pattern(p, style, offset)
        para.leading_wide = text[:leading].count("\u3000")
        if para.page_break_before or para.has_page_break:
            self.probe.page_breaks += 1
        self.probe.paragraphs.append(para)
        self._para_of[p] = para
        if where == "body":
            if self._last_table is not None:
                self._last_table.after_kind = _neighbor_kind(para)
                self._last_table = None
            self._prev_body = para
        return para

    def _bold_pattern(self, p, style, offset: int) -> tuple[str, bool]:
        """말머리 뒤 글자의 굵기 모양(전부·앞부분만·뒷부분만·섞임)과 밑줄 여부."""
        flags: list[bool] = []
        underlined = 0
        pos = 0
        for run in p.findall(qn("w:r")):
            text = _run_text(run)
            if not text:
                continue
            fmt = self._run_fmt(run, style)
            for ch in text:
                if pos >= offset and not ch.isspace():
                    flags.append(bool(fmt.bold))
                    underlined += 1 if fmt.underline else 0
                pos += 1
        if not flags:
            return "none", False
        if all(flags):
            pattern = "all"
        elif not any(flags):
            pattern = "none"
        else:
            changes = sum(1 for a, b in zip(flags, flags[1:]) if a != b)
            pattern = "prefix" if flags[0] and changes == 1 else "suffix" if flags[-1] and changes == 1 else "mixed"
        return pattern, underlined * 2 > len(flags)

    # ── 서식 풀기 ───────────────────────────────────────────────────────

    def _para_style(self, p):
        ppr = p.find(qn("w:pPr"))
        pstyle = ppr.find(qn("w:pStyle")) if ppr is not None else None
        if pstyle is not None and pstyle.get(qn("w:val")) in self.styles:
            return self.styles[pstyle.get(qn("w:val"))]
        return self.default_para_style

    def _chain(self, style):
        """기본 스타일 → … → 자기 순서(뒤로 갈수록 우선)."""
        out, seen = [], set()
        while style is not None and id(style) not in seen and len(out) < 12:
            seen.add(id(style))
            out.append(style)
            based = style.find(qn("w:basedOn"))
            style = self.styles.get(based.get(qn("w:val"))) if based is not None else None
        return list(reversed(out))

    def _effective_ppr(self, p, style) -> dict:
        """문단 속성의 실효값: docDefaults < 스타일 체인 < 번호 수준 < 직접 서식."""
        merged: dict = {}
        sources = []
        if self.defaults is not None:
            sources.append(self.defaults.find(f"{qn('w:pPrDefault')}/{qn('w:pPr')}"))
        sources += [s.find(qn("w:pPr")) for s in self._chain(style)]
        direct = p.find(qn("w:pPr"))
        lvl = self._level_def(*self._num_ref(p, style)) if self._num_ref(p, style) else None
        sources.append(lvl.find(qn("w:pPr")) if lvl is not None else None)
        sources.append(direct)
        for ppr in sources:
            if ppr is None:
                continue
            for tag in ("spacing", "ind"):
                el = ppr.find(qn(f"w:{tag}"))
                if el is not None:
                    merged.setdefault(tag, {}).update(_attrs(el))
            jc = ppr.find(qn("w:jc"))
            if jc is not None:
                merged["jc_value"] = jc.get(qn("w:val"))
            if ppr.find(qn("w:pageBreakBefore")) is not None:
                merged["pageBreakBefore"] = True
        # 들여쓰기에 hanging과 firstLine이 함께 오면 나중 것이 이긴다(스키마상 배타)
        ind = merged.get("ind", {})
        if "hanging" in ind and "firstLine" in ind and direct is not None:
            d_ind = direct.find(qn("w:ind"))
            if d_ind is not None:
                if d_ind.get(qn("w:hanging")) is not None:
                    ind.pop("firstLine", None)
                elif d_ind.get(qn("w:firstLine")) is not None:
                    ind.pop("hanging", None)
        return merged

    def _dominant_fmt(self, p, style) -> Fmt:
        """글자 수가 가장 많은 run의 실효 서식(제목 줄의 굵은 말머리 한 글자에 흔들리지 않게)."""
        best, best_len = None, -1
        for run in p.findall(qn("w:r")):
            length = len("".join(t.text or "" for t in run.findall(qn("w:t"))).strip())
            if length > best_len:
                best, best_len = run, length
        return self._run_fmt(best, style)

    def _run_fmt(self, run, pstyle) -> Fmt:
        rprs = []
        if self.defaults is not None:
            rprs.append(self.defaults.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}"))
        rprs += [s.find(qn("w:rPr")) for s in self._chain(pstyle)]
        if run is not None:
            direct = run.find(qn("w:rPr"))
            rstyle = direct.find(qn("w:rStyle")) if direct is not None else None
            if rstyle is not None and rstyle.get(qn("w:val")) in self.styles:
                rprs += [s.find(qn("w:rPr")) for s in self._chain(self.styles[rstyle.get(qn("w:val"))])]
            rprs.append(direct)
        fmt = Fmt()
        for rpr in rprs:
            if rpr is None:
                continue
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:
                fmt.east_asia = fonts.get(qn("w:eastAsia")) or fmt.east_asia
                fmt.ascii = fonts.get(qn("w:ascii")) or fmt.ascii
            sz = rpr.find(qn("w:sz"))
            if sz is not None and sz.get(qn("w:val")):
                fmt.size_pt = int(sz.get(qn("w:val"))) / 2
            for tag, attr in (("b", "bold"), ("i", "italic")):
                el = rpr.find(qn(f"w:{tag}"))
                if el is not None:
                    setattr(fmt, attr, el.get(qn("w:val")) not in ("0", "false", "off"))
            under = rpr.find(qn("w:u"))
            if under is not None:
                fmt.underline = under.get(qn("w:val")) != "none"
            scale = rpr.find(qn("w:w"))
            if scale is not None and scale.get(qn("w:val")):
                fmt.char_scale_pct = int(scale.get(qn("w:val")))
            color = rpr.find(qn("w:color"))
            if color is not None:
                value = color.get(qn("w:val"))
                theme = color.get(qn("w:themeColor"))
                if value and value != "auto":
                    fmt.color = value.upper()
                elif theme:
                    fmt.color = f"theme:{theme}"
                else:
                    fmt.color = None
        return fmt

    def _heading_level(self, p, style) -> int | None:
        for s in reversed(self._chain(style)):
            ppr = s.find(qn("w:pPr"))
            outline = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
            if outline is not None:
                value = int(outline.get(qn("w:val"), "9"))
                return value + 1 if value < 9 else None
        return None

    # ── 자동 번호 ───────────────────────────────────────────────────────

    def _num_ref(self, p, style) -> tuple[str, int] | None:
        ppr_list = [p.find(qn("w:pPr"))] + [s.find(qn("w:pPr")) for s in reversed(self._chain(style))]
        for ppr in ppr_list:
            numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
            if numpr is None:
                continue
            num_id = numpr.find(qn("w:numId"))
            ilvl = numpr.find(qn("w:ilvl"))
            value = num_id.get(qn("w:val")) if num_id is not None else "0"
            if value == "0":
                return None
            return value, int(ilvl.get(qn("w:val"))) if ilvl is not None else 0
        return None

    def _numbering_of(self, p, style) -> tuple[int, str | None, bool] | None:
        ref = self._num_ref(p, style)
        if ref is None:
            return None
        lvl = self._level_def(*ref)
        text = None
        pua = False
        if lvl is not None:
            el = lvl.find(qn("w:lvlText"))
            text = el.get(qn("w:val")) if el is not None else None
            pua = bool(text and _PUA.search(text))
        return ref[1], text, pua

    def _level_def(self, num_id: str, ilvl: int):
        if self.numbering is None:
            return None
        num = next((n for n in self.numbering.findall(qn("w:num"))
                    if n.get(qn("w:numId")) == num_id), None)
        if num is None:
            return None
        for override in num.findall(qn("w:lvlOverride")):
            if override.get(qn("w:ilvl")) == str(ilvl) and override.find(qn("w:lvl")) is not None:
                return override.find(qn("w:lvl"))
        abstract_id = num.find(qn("w:abstractNumId"))
        abstract = next(
            (a for a in self.numbering.findall(qn("w:abstractNum"))
             if abstract_id is not None and a.get(qn("w:abstractNumId")) == abstract_id.get(qn("w:val"))),
            None,
        )
        if abstract is None:
            return None
        return next((lvl for lvl in abstract.findall(qn("w:lvl"))
                     if lvl.get(qn("w:ilvl")) == str(ilvl)), None)

    # ── 표 ──────────────────────────────────────────────────────────────

    def _table(self, tbl, *, nested: bool) -> None:
        tblpr = tbl.find(qn("w:tblPr"))
        grid = tbl.find(qn("w:tblGrid"))
        widths = [
            round(_num(c.get(qn("w:w")), 0.0) / _TWIP_PER_MM, 1)
            for c in (grid.findall(qn("w:gridCol")) if grid is not None else [])
        ]
        rows = tbl.findall(qn("w:tr"))
        header_rows = 0
        for tr in rows:
            trpr = tr.find(qn("w:trPr"))
            if trpr is not None and trpr.find(qn("w:tblHeader")) is not None:
                header_rows += 1
            else:
                break
        merged = sum(
            1 for tc in tbl.iter(qn("w:tc"))
            if tc.find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}") is not None
            or tc.find(f"{qn('w:tcPr')}/{qn('w:vMerge')}") is not None
        )
        fill = None
        if rows:
            first_tc = rows[0].find(qn("w:tc"))
            shd = first_tc.find(f"{qn('w:tcPr')}/{qn('w:shd')}") if first_tc is not None else None
            fill = _shd_spec(shd)
        jc = tblpr.find(qn("w:jc")) if tblpr is not None else None
        margins: dict[str, float] = {}
        mar = tblpr.find(qn("w:tblCellMar")) if tblpr is not None else None
        if mar is not None:
            for side in ("top", "left", "bottom", "right"):
                el = mar.find(qn(f"w:{side}"))
                if el is None:
                    el = mar.find(qn("w:start" if side == "left" else "w:end"))
                if el is not None and el.get(qn("w:w")):
                    margins[side] = round(_num(el.get(qn("w:w")), 0.0) / _TWIP_PER_MM, 2)
        info = TableProbe(
            rows=len(rows), cols=len(widths), col_widths_mm=widths,
            jc=jc.get(qn("w:val")) if jc is not None else None,
            header_rows=header_rows, header_fill=fill, merged_cells=merged,
            cell_margin_mm=margins, nested=nested,
            total_width_mm=round(sum(widths), 1),
        )
        self._table_style(tblpr, info)
        self._table_fills(rows, info)
        if not nested:
            info.before_kind = (_neighbor_kind(self._prev_body) if self._prev_body is not None
                                else "(없음·표 바로 뒤)")
            self._prev_body = None
            self._last_table = info
        self.probe.tables.append(info)
        for r, tr in enumerate(rows):
            for c, tc in enumerate(tr.findall(qn("w:tc"))):
                start = len(self.probe.paragraphs)
                self._body(tc, where="table", nested=True)
                for para in self.probe.paragraphs[start:]:
                    if para.cell_row is None:   # 안쪽 표의 칸은 이미 자기 위치가 있다
                        para.cell_row, para.cell_col = r, c

    def _table_style(self, tblpr, info: TableProbe) -> None:
        if tblpr is None:
            return
        ind = tblpr.find(qn("w:tblInd"))
        if ind is not None:
            info.indent_mm = round(_num(ind.get(qn("w:w")), 0.0) / _TWIP_PER_MM, 1)
        look = tblpr.find(qn("w:tblLook"))
        if look is not None:
            first = look.get(qn("w:firstRow"))
            if first is not None:
                info.look_first_row = first in ("1", "true")
            elif look.get(qn("w:val")):
                info.look_first_row = bool(int(look.get(qn("w:val")), 16) & 0x0020)
        borders = tblpr.find(qn("w:tblBorders"))
        if borders is not None:
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = borders.find(qn(f"w:{side}"))
                if el is not None:
                    size = _num(el.get(qn("w:sz")), 0.0) / 8
                    info.borders[side] = f"{el.get(qn('w:val'))}/{size:g}pt"
        ref = tblpr.find(qn("w:tblStyle"))
        if ref is not None and ref.get(qn("w:val")) in self.styles:
            info.style_id = ref.get(qn("w:val"))
            for style in reversed(self._chain(self.styles[info.style_id])):
                for cond in style.findall(qn("w:tblStylePr")):
                    if cond.get(qn("w:type")) != "firstRow":
                        continue
                    spec = _shd_spec(cond.find(f"{qn('w:tcPr')}/{qn('w:shd')}"))
                    if spec:
                        info.style_first_row_fill = spec
                    bold = cond.find(f"{qn('w:rPr')}/{qn('w:b')}")
                    if bold is not None:
                        info.style_first_row_bold = bold.get(qn("w:val")) not in ("0", "false", "off")

    def _table_fills(self, rows, info: TableProbe) -> None:
        """칸 직접 음영이 첫 행에 있는가, 첫 열에 있는가, 그 밖에 있는가."""
        first_row = first_col = other = 0
        n_first_row = n_first_col = n_other = 0
        n_cells = bordered = 0
        colors: Counter = Counter()
        border_kinds: Counter = Counter()
        for r, tr in enumerate(rows):
            for c, tc in enumerate(tr.findall(qn("w:tc"))):
                fill = _shd_spec(tc.find(f"{qn('w:tcPr')}/{qn('w:shd')}"))
                shaded = fill is not None
                if shaded:
                    colors[fill] += 1
                borders = tc.find(f"{qn('w:tcPr')}/{qn('w:tcBorders')}")
                n_cells += 1
                if borders is not None:
                    kinds = {f"{el.get(qn('w:val'))}/{_num(el.get(qn('w:sz')), 0.0) / 8:g}pt"
                             for el in borders if el.get(qn("w:val")) not in (None, "nil", "none")}
                    if kinds:
                        bordered += 1
                        border_kinds.update(kinds)
                if r == 0:
                    n_first_row += 1
                    first_row += shaded
                elif c == 0:
                    n_first_col += 1
                    first_col += shaded
                else:
                    n_other += 1
                    other += shaded
        info.fill_first_row_share = round(first_row / n_first_row, 2) if n_first_row else None
        info.fill_first_col_share = round(first_col / n_first_col, 2) if n_first_col else None
        info.fill_other_share = round(other / n_other, 2) if n_other else None
        info.fill_colors = [k for k, _ in colors.most_common(3)]
        info.cell_border_share = round(bordered / n_cells, 2) if n_cells else None
        info.cell_border_kinds = [k for k, _ in border_kinds.most_common(3)]

    # ── 쪽·머리말·텍스트 상자 ───────────────────────────────────────────

    def _sections(self) -> None:
        for index, section in enumerate(self.docx.sections):
            width, height = section.page_width, section.page_height
            self.probe.sections.append(PageProbe(
                width_mm=_mm(width), height_mm=_mm(height),
                top_mm=_mm(section.top_margin), bottom_mm=_mm(section.bottom_margin),
                left_mm=_mm(section.left_margin), right_mm=_mm(section.right_margin),
                header_mm=_mm(section.header_distance), footer_mm=_mm(section.footer_distance),
                orientation="landscape" if (width or 0) > (height or 0) else "portrait",
                paper=_paper_name(width, height),
                title_page=bool(section.different_first_page_header_footer),
            ))
            for kind, part in (("header", section.header), ("footer", section.footer)):
                element = part._element  # noqa: SLF001 — 머리말 XML을 직접 본다
                text = "\n".join(_text(p) for p in element.iter(qn("w:p")) if _text(p).strip())
                first = next(element.iter(qn("w:p")), None)
                jc = first.find(f"{qn('w:pPr')}/{qn('w:jc')}") if first is not None else None
                self.probe.headers_footers.append(HeaderFooterProbe(
                    kind=kind, section=index, text_len=len(text.strip()), text=text,
                    has_page_field=_has_page_field(element),
                    align=jc.get(qn("w:val")) if jc is not None else None,
                    linked=bool(part.is_linked_to_previous),
                ))

    def _textboxes(self, body) -> None:
        for box in body.iter(qn("w:txbxContent")):
            if _inside_fallback(box):
                continue  # mc:AlternateContent의 Fallback은 같은 상자의 사본
            self.probe.textbox_count += 1
            index = len(self.probe.textboxes)
            made = []
            for p in box.findall(qn("w:p")):
                para = self._paragraph(p, "textbox")
                para.box = index
                made.append(para)
            info = _box_geometry(box)
            info.paragraphs = sum(1 for x in made if not x.blank)
            info.text_len = sum(x.text_len for x in made if not x.blank)
            anchor = _anchor_paragraph(box)
            anchor_para = self._para_of.get(anchor) if anchor is not None else None
            if anchor_para is not None:
                info.anchor_where = anchor_para.where
                info.anchor_marker = anchor_para.marker
                info.anchor_text_len = anchor_para.text_len
            info.placement = _placement(info, self.probe.sections[0] if self.probe.sections else None)
            self.probe.textboxes.append(info)

    def _has_toc(self, body) -> bool:
        for sdt in body.iter(qn("w:sdt")):
            gallery = sdt.find(f"{qn('w:sdtPr')}/{qn('w:docPartObj')}/{qn('w:docPartGallery')}")
            if gallery is not None and "table of contents" in (gallery.get(qn("w:val")) or "").lower():
                return True
        return any("TOC" in (t.text or "").split()[:1] or (t.text or "").strip().startswith("TOC")
                   for t in body.iter(qn("w:instrText")))


# ── 도우미 ──────────────────────────────────────────────────────────────


def _text(el) -> str:
    """문단의 보이는 글자. 지워진 변경 내용·필드 코드·텍스트 상자 속 글자는 뺀다."""
    out: list[str] = []

    def walk(node) -> None:
        for child in node:
            tag = child.tag
            if tag in (qn("w:del"), qn("w:instrText"), qn("w:txbxContent"), _MC_FALLBACK):
                continue
            if tag == qn("w:t"):
                out.append(child.text or "")
            elif tag == qn("w:tab"):
                out.append("\t")
            elif tag in (qn("w:br"), qn("w:cr")) and child.get(qn("w:type")) != "page":
                out.append("\n")
            else:
                walk(child)

    walk(el)
    return "".join(out)


def _typed_marker(text: str) -> tuple[str, str] | None:
    """(말머리, 뒤 구분: tab|space|none). 말머리가 아니면 None."""
    match = _TYPED_MARKER.match(text)
    if match is None or is_blank(text):
        return None
    mark = match.group("sym") or match.group("wordy")
    raw = match.group("sep")
    sep = "tab" if raw == "\t" else "space" if raw else "none"
    if match.group("wordy") and sep == "none":
        return None
    rest = text[match.end("sym") if match.group("sym") else match.end("wordy"):]
    if rest[:1] == mark[-1:]:
        return None  # "○○팀"처럼 같은 기호가 연달아 오면 자리표시자
    return mark, sep


def _has_page_break(p) -> bool:
    return any(br.get(qn("w:type")) == "page" for br in p.iter(qn("w:br")))


def _has_page_field(element) -> bool:
    for instr in element.iter(qn("w:instrText")):
        if (instr.text or "").strip().upper().startswith(("PAGE", "NUMPAGES")):
            return True
    return any((f.get(qn("w:instr")) or "").strip().upper().startswith(("PAGE", "NUMPAGES"))
               for f in element.iter(qn("w:fldSimple")))


def _shd_spec(shd) -> str | None:
    """음영 요소를 한 줄 설명으로: 직접 색(RRGGBB) | theme:이름+농도 | pattern:모양/색. 없으면 None."""
    if shd is None:
        return None
    fill, theme, val = shd.get(qn("w:fill")), shd.get(qn("w:themeFill")), shd.get(qn("w:val"))
    if theme:
        shade, tint = shd.get(qn("w:themeFillShade")), shd.get(qn("w:themeFillTint"))
        return f"theme:{theme}" + (f"+shade{shade}" if shade else "") + (f"+tint{tint}" if tint else "")
    if fill and fill.lower() not in ("auto", "ffffff"):
        return fill.upper()
    if val and val not in ("clear", "nil", "none"):
        color = shd.get(qn("w:color"))
        return f"pattern:{val}/{color or 'auto'}"
    return None


def _run_text(run) -> str:
    out: list[str] = []
    for child in run:
        if child.tag == qn("w:t"):
            out.append(child.text or "")
        elif child.tag == qn("w:tab"):
            out.append("\t")
        elif child.tag in (qn("w:br"), qn("w:cr")) and child.get(qn("w:type")) != "page":
            out.append("\n")
    return "".join(out)


_OPENERS = "[［【〔〈《「『"


def _neighbor_kind(para: ParaProbe) -> str:
    """표 바로 위·아래 줄이 어떤 줄인가 (표 제목·주석 관례를 본다)."""
    text = para.text.strip()
    if para.blank:
        return "빈 줄"
    if text.startswith(("<표", "〈표", "표 ")):
        return "표 번호"
    if text[:1] and text[0] in _OPENERS:
        return "꺾쇠 제목"
    if para.marker == "※" or text.startswith("*"):
        return "※ 주석"
    if para.marker_kind:
        return f"말머리 {para.marker}"
    return "문단(말머리 없음)"


def _inside_fallback(el) -> bool:
    parent = el.getparent()
    while parent is not None:
        if parent.tag == _MC_FALLBACK:
            return True
        parent = parent.getparent()
    return False


def _box_shape(box):
    """txbxContent를 감싼 도형 요소(DrawingML wp:anchor·wp:inline 또는 VML v:shape 등)."""
    node = box.getparent()
    while node is not None and isinstance(node.tag, str):
        if node.tag in (_WP + "anchor", _WP + "inline"):
            return node
        if node.tag.startswith(_VML) and node.tag[len(_VML):] in _VML_SHAPES:
            return node
        if node.tag == qn("w:body"):
            break
        node = node.getparent()
    return None


def _anchor_paragraph(box):
    """상자가 붙어 있는 본문·표 문단(상자 안의 문단이 아니라 상자를 품은 바깥 문단)."""
    shape = _box_shape(box)
    node = shape.getparent() if shape is not None else box.getparent()
    while node is not None and isinstance(node.tag, str):
        if node.tag == qn("w:p"):
            return node
        node = node.getparent()
    return None


def _css_mm(value: str | None) -> float | None:
    """VML style의 길이("12pt", "3cm")를 mm로. 단위가 없으면 px로 본다."""
    if not value:
        return None
    m = re.fullmatch(r"\s*(-?[0-9.]+)\s*(pt|in|cm|mm|px)?\s*", value)
    if not m:
        return None
    number = float(m.group(1))
    factor = {"pt": 25.4 / 72, "in": 25.4, "cm": 10.0, "mm": 1.0, "px": 25.4 / 96}[m.group(2) or "px"]
    return round(number * factor, 1)


def _box_geometry(box) -> TextBoxProbe:
    shape = _box_shape(box)
    if shape is None:
        return TextBoxProbe(kind="unknown", floating=False)
    if shape.tag.startswith(_WP):
        return _drawingml_geometry(shape)
    return _vml_geometry(shape)


def _emu_mm(text: str | None) -> float | None:
    try:
        return round(int(text) / EMU_PER_MM, 1) if text else None
    except ValueError:
        return None


def _drawingml_geometry(shape) -> TextBoxProbe:
    info = TextBoxProbe(kind="drawingml", floating=shape.tag == _WP + "anchor")
    for axis in ("H", "V"):
        pos = shape.find(f"{_WP}position{axis}")
        if pos is None:
            continue
        off, align = pos.find(_WP + "posOffset"), pos.find(_WP + "align")
        key = axis.lower()
        setattr(info, f"{key}_rel", pos.get("relativeFrom"))
        setattr(info, f"{key}_offset_mm", _emu_mm(off.text) if off is not None else None)
        setattr(info, f"{key}_align", align.text if align is not None else None)
    extent = shape.find(_WP + "extent")
    if extent is not None:
        info.width_mm, info.height_mm = _emu_mm(extent.get("cx")), _emu_mm(extent.get("cy"))
    info.behind = shape.get("behindDoc") in ("1", "true") if shape.get("behindDoc") is not None else None
    for child in shape:
        if isinstance(child.tag, str) and child.tag.startswith(_WP + "wrap"):
            info.wrap = "wrap" + child.tag[len(_WP + "wrap"):]
    sppr = next(shape.iter(_WPS + "spPr"), None)
    if sppr is not None:
        line = sppr.find(_A + "ln")
        if line is not None:
            if line.find(_A + "noFill") is not None:
                info.border = "none"
            else:
                info.border = "line"
                color = line.find(f"{_A}solidFill/{_A}srgbClr")
                info.border_color = color.get("val").upper() if color is not None and color.get("val") else None
        if sppr.find(_A + "noFill") is not None:
            info.fill = "none"
        else:
            color = sppr.find(f"{_A}solidFill/{_A}srgbClr")
            if color is not None and color.get("val"):
                info.fill = color.get("val").upper()
    return info


def _vml_geometry(shape) -> TextBoxProbe:
    style = {}
    for part in (shape.get("style") or "").split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            style[k.strip()] = v.strip()
    info = TextBoxProbe(kind="vml", floating=style.get("position") == "absolute")
    info.h_rel = style.get("mso-position-horizontal-relative")
    info.v_rel = style.get("mso-position-vertical-relative")
    info.h_align = style.get("mso-position-horizontal")
    info.v_align = style.get("mso-position-vertical")
    info.h_offset_mm = _css_mm(style.get("margin-left", style.get("left")))
    info.v_offset_mm = _css_mm(style.get("margin-top", style.get("top")))
    info.width_mm, info.height_mm = _css_mm(style.get("width")), _css_mm(style.get("height"))
    wrap = shape.find("{urn:schemas-microsoft-com:office:word}wrap")
    info.wrap = wrap.get("type") if wrap is not None else None
    if (shape.get("stroked") or "").lower() in ("f", "false"):
        info.border = "none"
    elif shape.get("stroked") is not None or shape.get("strokecolor"):
        info.border = "line"
        info.border_color = (shape.get("strokecolor") or "").lstrip("#").upper() or None
    if (shape.get("filled") or "").lower() in ("f", "false"):
        info.fill = "none"
    elif shape.get("fillcolor"):
        info.fill = shape.get("fillcolor").lstrip("#").upper()
    return info


def _placement(info: TextBoxProbe, page: PageProbe | None) -> str:
    """'가로 구역 / 세로 기준' 한 줄 요약. 본문 문장의 옆 여백인지 아래인지를 가른다."""
    if not info.floating:
        return "줄 안(인라인)"
    horizontal = "가로 불명"
    if info.h_align:
        horizontal = f"가로 {info.h_align}"
    elif info.h_offset_mm is not None and page is not None:
        start = {"page": 0.0, "margin": page.left_mm, "column": page.left_mm, "leftMargin": 0.0,
                 "rightMargin": page.width_mm - page.right_mm, "character": page.left_mm,
                 "insideMargin": page.left_mm, "outsideMargin": page.left_mm}.get(info.h_rel or "")
        if start is not None:
            x = start + info.h_offset_mm
            info.x_mm = round(x, 1)
            width = info.width_mm or 0.0
            if x >= page.width_mm - page.right_mm - 1:
                horizontal = "오른쪽 여백"
            elif x + width <= page.left_mm + 1:
                horizontal = "왼쪽 여백"
            else:
                horizontal = "본문 영역 안"
    vertical = "세로 불명"
    if info.v_align:
        vertical = f"세로 {info.v_align}"
    elif info.v_offset_mm is not None:
        if info.v_rel in ("paragraph", "line", "text"):
            vertical = "문단 기준 아래" if info.v_offset_mm >= 0 else "문단 기준 위"
        else:
            vertical = f"{info.v_rel or '쪽'} 기준"
    return f"{horizontal} / {vertical}"


def _attrs(el) -> dict[str, str]:
    return {k.replace(_W, ""): v for k, v in el.attrib.items()}


def _num(value, default: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _mm(length) -> float:
    return round((length or 0) / EMU_PER_MM, 1)


def _paper_name(width, height) -> str | None:
    for name, (w, h) in PAGE_SIZES.items():
        for a, b in ((w, h), (h, w)):
            if abs((width or 0) - parse_length(a)) < 1.5 * EMU_PER_MM and \
                    abs((height or 0) - parse_length(b)) < 1.5 * EMU_PER_MM:
                return name
    return None


def _style_name(style) -> str | None:
    if style is None:
        return None
    name = style.find(qn("w:name"))
    return name.get(qn("w:val")) if name is not None else style.get(qn("w:styleId"))
