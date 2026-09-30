"""설치·환경 진단 — `doc2report doctor` / `doctor.bat`.

2026-09-30 사용자: 외부에서 만든 코드를 사내 PC에 설치하느라 고생이 컸다(파이썬 환경·uv·프록시·인증서).
팀원이 같은 고생을 하지 않게, 막히는 지점을 한 번에 점검하고 **무엇을 하면 되는지**까지 알려 준다.
결과는 파일로도 남겨 담당자에게 그대로 보낼 수 있게 한다(토큰 값은 절대 쓰지 않는다).

점검 항목은 실제로 겪은 문제에서 왔다(CLAUDE.md "겪은 함정"·Confluence/온프렘 절):
- 파이썬·패키지가 다 들어왔는지, 어느 코드(브랜치·커밋)를 돌리는지(옛 브랜치에 남아 "고친 게 반영 안 됨")
- Word 기본 템플릿이 열리는지(사내 문서보안이 .docx를 손상시킨 일), 결과 폴더에 쓸 수 있는지
- 표 폭 계산용 글꼴(서식 프로파일에 적힌 것)을 찾는지, PDF 변환기(Word) 유무
- 사용자 등록·토큰 암호화(DPAPI)가 되는지
- 프록시 환경변수가 사내 LLM 요청을 가로채는지(HTTP_PROXY → Squid 403), Confluence·LLM에 실제로 닿는지
"""

from __future__ import annotations

import os
import platform
import socket
import sys
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parents[2]  # 저장소(또는 설치 묶음) 맨 위

OK, WARN, FAIL, INFO = "OK", "주의", "실패", "참고"


@dataclass
class Check:
    level: str
    title: str
    detail: str = ""
    fix: str = ""


def run(*, network: bool = True) -> list[Check]:
    checks: list[Check] = []
    for step in (_python, _packages, _code, _template, _output_dir, _fonts, _pdf, _account, _proxy):
        checks.extend(_safe(step))
    if network:
        checks.extend(_safe(_confluence))
        checks.extend(_safe(_llm))
    checks.extend(_safe(_port))
    return checks


def report(checks: list[Check]) -> str:
    lines = [f"doc2report 진단 결과 — {datetime.now():%Y-%m-%d %H:%M}",
             f"위치: {ROOT}", ""]
    for check in checks:
        lines.append(f"[{check.level}] {check.title}")
        if check.detail:
            lines.append(f"      {check.detail}")
        if check.fix:
            lines.append(f"      → 해결: {check.fix}")
    fails = sum(c.level == FAIL for c in checks)
    warns = sum(c.level == WARN for c in checks)
    lines += ["", f"요약: 실패 {fails}건, 주의 {warns}건"
              + (" — 그대로 쓰면 됩니다." if not fails else " — '해결'에 적힌 대로 한 뒤 다시 진단하세요.")]
    return "\n".join(lines)


def _safe(step) -> list[Check]:
    try:
        result = step()
        return result if isinstance(result, list) else [result]
    except Exception as exc:  # 진단이 진단 도중 죽으면 안 된다
        return [Check(FAIL, f"{step.__name__.strip('_')} 점검 중 오류", f"{type(exc).__name__}: {exc}")]


# ── 점검 ────────────────────────────────────────────────────────────────


def _python() -> Check:
    runtime = ROOT / "runtime"
    where = "설치 묶음(runtime)" if Path(sys.executable).parent == runtime else sys.executable
    ok = sys.version_info >= (3, 11)
    return Check(OK if ok else FAIL, f"파이썬 {platform.python_version()} ({platform.architecture()[0]})",
                 f"실행 위치: {where}",
                 "" if ok else "파이썬 3.11 이상이 필요합니다 — 설치 묶음(zip)을 쓰면 파이썬을 따로 깔 필요가 없습니다")


def _packages() -> Check:
    import importlib.metadata as md

    need = ["python-docx", "lxml", "pydantic", "httpx", "PyYAML", "fonttools", "typer",
            "markdown-it-py", "mdit-py-plugins"]
    missing, found = [], []
    for name in need:
        try:
            found.append(f"{name} {md.version(name)}")
        except md.PackageNotFoundError:
            missing.append(name)
    if missing:
        return Check(FAIL, "필요한 패키지가 빠짐", ", ".join(missing),
                     "설치 묶음(zip)을 새로 풀거나, 개발 환경이면 wheelhouse로 다시 설치하세요(docs/team-setup.md)")
    return Check(OK, "필요한 패키지 모두 있음", ", ".join(found))


def _code() -> Check:
    build = ROOT / "runtime" / "BUILD.txt"
    head = ROOT / ".git" / "HEAD"
    if head.is_file():
        ref = head.read_text(encoding="utf-8").strip()
        branch = ref.rsplit("/", 1)[-1] if ref.startswith("ref:") else "(브랜치 아님)"
        commit = _git_commit(ref)
        if branch != "main":
            return Check(WARN, f"코드: git {branch} 브랜치 {commit}",
                         "작업 브랜치는 main 하나입니다 — 다른 브랜치면 git pull 해도 새 코드가 안 들어옵니다",
                         "git checkout main 후 git pull origin main")
        return Check(OK, f"코드: git main {commit}")
    if build.is_file():
        return Check(OK, "코드: 설치 묶음", build.read_text(encoding="utf-8").strip().replace("\n", " · "))
    return Check(INFO, "코드: 버전 정보 없음")


def _git_commit(ref: str) -> str:
    if not ref.startswith("ref:"):
        return ref[:8]
    path = ROOT / ".git" / ref[5:].strip()
    if path.is_file():
        return path.read_text(encoding="utf-8").strip()[:8]
    packed = ROOT / ".git" / "packed-refs"
    if packed.is_file():
        for line in packed.read_text(encoding="utf-8").splitlines():
            if line.endswith(ref[5:].strip()):
                return line.split()[0][:8]
    return ""


def _template() -> Check:
    from docx import Document

    from .render.base_template import open_base_template

    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / "진단.docx"
        doc = Document(open_base_template())
        doc.add_paragraph("진단")
        doc.save(str(path))
        Document(str(path))  # 다시 열리는지 — 문서보안이 파일을 바꾸면 여기서 실패
    return Check(OK, "Word 문서 만들고 다시 열기")


def _output_dir() -> Check:
    folder = Path(os.environ.get("DOC2REPORT_OUTPUT_DIR") or (Path.cwd() / "out" / "webapp"))
    try:
        folder.mkdir(parents=True, exist_ok=True)
        probe = folder / ".doc2report-write-test"
        probe.write_text("x", encoding="utf-8")
        probe.unlink()
    except OSError as exc:
        return Check(FAIL, "결과 저장 폴더에 쓸 수 없음", f"{folder}: {exc}",
                     "다른 폴더를 쓰세요: start_webapp.bat --output-dir D:\\보고서")
    return Check(OK, "결과 저장 폴더", str(folder))


def _fonts() -> list[Check]:
    """서식(preset) 프로파일들이 쓰는 본문·표 글꼴 — 글꼴 이름은 코드가 아니라 프로파일에만 있다."""
    from .layout.measure import get_metrics
    from .profile import PROFILE_DIR, load_profile

    names: list[str] = []
    for path in sorted(PROFILE_DIR.glob("*.yaml")):
        try:
            prof = load_profile(path.stem)
        except Exception:
            continue
        for key in ("body", "table"):
            name = prof.font(key).east_asia
            if name and name not in names:
                names.append(name)
    out = []
    for name in names:
        loaded = get_metrics(name).loaded
        out.append(Check(OK if loaded else WARN, f"글꼴 {name}",
                         "" if loaded else "글꼴 파일을 못 찾아 표 폭을 근사값으로 계산합니다(표가 약간 어긋날 수 있음)",
                         "" if loaded or platform.system() != "Windows" else
                         "Windows 글꼴 폴더에 해당 글꼴이 있는지 확인하세요"))
    return out


def _pdf() -> Check:
    from .render.pdf import pdf_engine

    engine = pdf_engine()
    if engine["available"]:
        return Check(OK, f"PDF 변환: {engine['engine']}")
    return Check(WARN, "PDF 변환기 없음", "Word 문서(.docx)는 만들 수 있고 PDF만 안 됩니다",
                 "MS Word가 설치된 PC에서 쓰세요")


def _account() -> list[Check]:
    from . import account

    out = []
    try:
        registered = account.load()
        if registered:
            account.apply(registered)  # 뒤의 연결 점검이 등록한 토큰·주소로 하도록
    except Exception as exc:
        return [Check(FAIL, "사용자 등록 정보를 읽지 못함", str(exc),
                      "다른 PC·다른 Windows 계정에서 만든 등록이면 화면에서 '등록 삭제' 후 다시 등록하세요")]
    if registered is None:
        out.append(Check(WARN, "사용자 미등록", f"보관 위치: {account.path()}",
                         "start_webapp.bat으로 화면을 열면 등록 창이 뜹니다"))
    else:
        who = registered.verified_as or "토큰 주인 확인 안 됨"
        out.append(Check(OK if registered.verified_as else WARN, f"사용자 등록: {registered.name}",
                         f"Confluence 토큰 주인: {who} · 보관: {account.protection()}",
                         "" if registered.verified_as else "화면 위 사용자 이름 → '토큰 주인 다시 확인'"))
    if platform.system() == "Windows":
        sealed = account._seal("doc2report-self-test")
        ok = account._unseal(sealed) == "doc2report-self-test"
        out.append(Check(OK if ok else FAIL, "토큰 암호화(DPAPI)"))
    return out


def _proxy() -> Check:
    proxy = os.environ.get("HTTPS_PROXY") or os.environ.get("HTTP_PROXY") or \
        os.environ.get("https_proxy") or os.environ.get("http_proxy")
    if not proxy:
        return Check(OK, "프록시 환경변수 없음")
    no_proxy = (os.environ.get("NO_PROXY") or os.environ.get("no_proxy") or "").lower()
    missing = [host for host in (_host("DOC2REPORT_LLM_BASE_URL"), _host("CONFLUENCE_URL"))
               if host and host.lower() not in no_proxy]
    shown = no_proxy if len(no_proxy) <= 120 else no_proxy[:117] + "…"
    detail = f"HTTP(S)_PROXY={_mask_proxy(proxy)} · NO_PROXY={shown or '(없음)'}"
    if missing:
        return Check(WARN, "사내 주소가 프록시를 거칠 수 있음", detail + f" · 예외에 없는 주소: {', '.join(missing)}",
                     "사내 LLM은 등록된 주소를 자동으로 프록시 예외에 넣습니다. 그래도 403이 나면 "
                     "네트워크 담당자에게 프록시 예외 등록을 요청하세요")
    return Check(OK, "프록시 예외 설정", detail)


def _confluence() -> Check:
    if not (os.environ.get("CONFLUENCE_URL") and os.environ.get("CONFLUENCE_API_TOKEN")):
        return Check(INFO, "Confluence: 설정 없음(연결 점검 건너뜀)", "",
                     "Confluence를 쓰려면 화면의 사용자 등록에서 주소·개인 토큰을 넣으세요")
    from .sources.confluence import confluence_whoami

    try:
        who = confluence_whoami()
    except Exception as exc:
        text = str(exc)
        fix = ("사내 루트 인증서 문제일 수 있습니다 — Windows에서는 자동으로 PowerShell로 다시 시도합니다. "
               "계속 실패하면 담당자에게 이 진단 파일을 보내 주세요")
        if "401" in text or "익명" in text:
            fix = "토큰이 틀렸거나 만료됐습니다 — Confluence에서 새로 발급해 다시 등록하세요"
        return Check(FAIL, "Confluence 연결 실패", text[:300], fix)
    return Check(OK, "Confluence 연결", f"토큰 주인: {who}")


def _llm() -> Check:
    base = os.environ.get("DOC2REPORT_LLM_BASE_URL")
    if not base:
        return Check(INFO, "온프렘 LLM: 설정 없음(연결 점검 건너뜀)")
    import httpx

    key = os.environ.get("DOC2REPORT_LLM_API_KEY")
    headers = {"Authorization": f"Bearer {key}"} if key else {}
    try:
        resp = httpx.get(base.rstrip("/") + "/models", headers=headers, timeout=10)
    except Exception as exc:
        return Check(FAIL, "온프렘 LLM에 연결할 수 없음", f"{base}: {exc}",
                     "주소·포트를 확인하세요. 사내 프록시가 막는 경우가 있었습니다(진단의 프록시 항목 참고)")
    if resp.status_code == 403 and "squid" in resp.text.lower():
        return Check(FAIL, "온프렘 LLM 요청이 사내 프록시에서 막힘(403)", base,
                     "NO_PROXY에 LLM 서버 주소를 넣거나 네트워크 담당자에게 예외 등록을 요청하세요")
    if resp.status_code >= 400:
        return Check(WARN, f"온프렘 LLM 응답 HTTP {resp.status_code}", resp.text[:200],
                     "주소 끝이 /v1 인지, API 키가 필요한 서버인지 확인하세요")
    model = os.environ.get("DOC2REPORT_MODEL", "")
    ids = [m.get("id") for m in (resp.json().get("data") or []) if isinstance(m, dict)]
    if model and ids and model not in ids:
        return Check(WARN, "온프렘 LLM: 등록한 모델명이 서버 목록에 없음", f"서버 모델: {', '.join(ids)}",
                     f"사용자 등록의 모델명을 서버 목록 중 하나로 고치세요(지금: {model})")
    return Check(OK, f"온프렘 LLM 연결: {model}", base)


def _port(port: int = 8765) -> Check:
    with socket.socket() as sock:
        sock.settimeout(0.5)
        busy = sock.connect_ex(("127.0.0.1", port)) == 0
    if busy:
        return Check(INFO, f"포트 {port}를 쓰는 프로그램이 있음",
                     "doc2report가 이미 켜져 있을 수 있습니다(새로 켜면 자동으로 넘겨받음)")
    return Check(OK, f"포트 {port} 비어 있음")


def _host(env: str) -> str:
    return urlparse(os.environ.get(env) or "").hostname or ""


def _mask_proxy(url: str) -> str:
    parsed = urlparse(url if "://" in url else "http://" + url)
    return f"{parsed.hostname}:{parsed.port}" if parsed.hostname else url
